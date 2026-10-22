"""Autorización HTTP con sesiones y PostgreSQL reales, sobre una base desechable.

Ejecutar con ``python tests/postgres/ejecutar.py --patron test_autorizacion_pg.py``.
Se conservan el login y CSRF reales: las solicitudes directas no omiten las
protecciones que utiliza el navegador.
"""

import re
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from html import unescape
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import text

from entorno_pg import CasoPostgres, HASH_PRUEBA


MENSAJE_403 = 'No tienes permiso para realizar esta acción'


class AutorizacionPostgres(CasoPostgres):
    def _valor(self, sql, **datos):
        with self.engine.connect() as conexion:
            return conexion.execute(text(sql), datos).scalar_one()

    def _fila(self, sql, **datos):
        with self.engine.connect() as conexion:
            return conexion.execute(text(sql), datos).mappings().one()

    def _sql(self, sql, **datos):
        with self.engine.begin() as conexion:
            return conexion.execute(text(sql), datos)

    def _crear(self, sql, **datos):
        with self.engine.begin() as conexion:
            return conexion.execute(text(sql), datos).scalar_one()

    def _csrf(self, cliente=None, pagina='/productos'):
        cliente = cliente or self.client
        respuesta = cliente.get(pagina)
        self.assertEqual(respuesta.status_code, 200, pagina)
        html = respuesta.get_data(as_text=True)
        token = re.search(r'name="csrf_token"\s+value="([^"]+)"', html)
        if not token:
            token = re.search(r'data-csrf-token="([^"]+)"', html)
        self.assertIsNotNone(token, f'No se encontró CSRF en {pagina}')
        return unescape(token.group(1))

    def _login(self, nombre='admin', *, cliente=None):
        cliente = cliente or self.app.app.test_client()
        respuesta = cliente.post('/login', data={
            'correo': f'{nombre}@example.test',
            'password': 'PruebaPostgres2026!',
            'csrf_token': self._csrf(cliente, '/login'),
        })
        self.assertEqual(respuesta.status_code, 302)
        destino = {'admin': '/', 'cajero1': '/pos', 'cajero2': '/pos',
                   'bodeguero': '/inventario', 'otroadmin': '/'}[nombre]
        self.assertEqual(respuesta.location, destino)
        self.client = cliente
        return cliente

    def _form(self, ruta, datos=None, *, cliente=None, token=None):
        cliente = cliente or self.client
        return cliente.post(ruta, data={
            **(datos or {}), 'csrf_token': token or self._csrf(cliente),
        })

    def _denegado(self, respuesta):
        self.assertEqual(respuesta.status_code, 403)
        if respuesta.is_json:
            self.assertIn(MENSAJE_403, str(respuesta.get_json()))
        else:
            self.assertIn(MENSAJE_403, respuesta.get_data(as_text=True))

    def _stock(self, producto=1):
        return self._valor('SELECT existencia_actual FROM vista_inventario WHERE id_producto=:id',
                           id=producto)

    def _estado_ventas(self):
        return tuple(self._valor(f'SELECT COUNT(*) FROM {tabla}') for tabla in (
            'venta', 'detalle_venta', 'movimiento_inventario', 'movimiento_caja'))

    def _payload(self, **cambios):
        return {
            'cliente_id': 1, 'metodo': 'EFECTIVO', 'recibido': '100',
            'items': [{'id': 1, 'cantidad': 1, 'descuento': '0'}], **cambios,
        }

    def _cobrar(self, *, cliente=None, **cambios):
        cliente = cliente or self.client
        return cliente.post('/pos/cobrar', json=self._payload(**cambios),
                            headers={'X-CSRFToken': self._csrf(cliente), 'Idempotency-Key': uuid4().hex})

    def _venta(self, cajero='cajero1', *, producto=1, cantidad=1, cliente_id=1):
        self._login(cajero)
        respuesta = self._cobrar(cliente_id=cliente_id,
                                items=[{'id': producto, 'cantidad': cantidad, 'descuento': '0'}])
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.get_json()['ok'])
        return respuesta.get_json()

    def _otro_admin(self):
        return self._crear('''INSERT INTO usuario(id_rol,nombre,correo,hash_contrasena)
            VALUES(1,'Otro Administrador','otroadmin@example.test',:hash) RETURNING id_usuario''',
                           hash=HASH_PRUEBA)

    def test_sin_sesion_las_rutas_protegidas_exigen_autenticacion(self):
        for ruta in ('/', '/productos', '/inventario', '/pos', '/ventas',
                     '/ventas/1/ticket', '/gestion/clientes', '/gestion/proveedores',
                     '/reportes', '/reportes/excel', '/usuarios', '/configuracion'):
            with self.subTest(ruta=ruta):
                respuesta = self.client.get(ruta)
                self.assertEqual(respuesta.status_code, 302)
                self.assertIn('/login', respuesta.location)

    def test_login_redirige_por_rol_y_no_duplica_los_roles_existentes(self):
        for perfil in ('admin', 'cajero1', 'bodeguero'):
            with self.subTest(perfil=perfil):
                self._login(perfil)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM rol'), 3)
        for nombre in ('Administrador', 'Cajero', 'Bodeguero'):
            self.assertEqual(self._valor('SELECT COUNT(*) FROM rol WHERE nombre=:nombre',
                                         nombre=nombre), 1)

    def test_matriz_consultas_mediante_urls_directas(self):
        rutas = {
            '/productos': {'admin', 'cajero1', 'bodeguero'},
            '/productos/nuevo': {'admin'}, '/productos/1/editar': {'admin'},
            '/gestion/categorias': {'admin'},
            '/gestion/clientes': {'admin', 'cajero1'},
            '/gestion/proveedores': {'admin', 'bodeguero'},
            '/pos': {'admin', 'cajero1'}, '/caja': {'admin', 'cajero1'},
            '/ventas': {'admin', 'cajero1'}, '/reportes': {'admin', 'cajero1'},
            '/reportes/excel': {'admin', 'cajero1'},
            '/reportes/ventas/excel': {'admin', 'cajero1'},
            '/compras': {'admin'}, '/inventario': {'admin', 'bodeguero'},
            '/inventario/excel': {'admin', 'bodeguero'},
            '/inventario/ajustes': {'admin', 'bodeguero'},
            '/usuarios': {'admin'}, '/configuracion': {'admin'},
        }
        for perfil in ('admin', 'cajero1', 'bodeguero'):
            self._login(perfil)
            for ruta, permitidos in rutas.items():
                with self.subTest(perfil=perfil, ruta=ruta):
                    respuesta = self.client.get(ruta)
                    if perfil in permitidos:
                        self.assertEqual(respuesta.status_code, 200)
                    else:
                        self._denegado(respuesta)

    def test_mutaciones_administrativas_no_aceptan_cajero_ni_bodeguero(self):
        casos = (
            ('/productos/nuevo', {}), ('/productos/1/editar', {'costo': '1', 'precio_venta': '2'}),
            ('/productos/1/estado', {}), ('/productos/1/eliminar', {}),
            ('/gestion/categorias', {'nombre': 'No autorizado'}),
            ('/gestion/categorias/1/estado', {'estado': 'desactivar'}),
            ('/gestion/categorias/1/eliminar', {}),
            ('/gestion/clientes', {'id': '1', 'nombre_razon_social': 'Modificado'}),
            ('/gestion/clientes/1/estado', {'estado': 'desactivar'}),
            ('/gestion/clientes/1/eliminar', {}),
            ('/gestion/proveedores', {'razon_social': 'No autorizado'}),
            ('/gestion/proveedores', {'id': '1', 'razon_social': 'Modificado'}),
            ('/gestion/proveedores/1/estado', {'estado': 'desactivar'}),
            ('/gestion/proveedores/1/eliminar', {}),
            ('/compras/nueva', {}), ('/usuarios', {}),
            ('/usuarios/1/estado', {'estado': 'desactivar'}),
            ('/usuarios/1/eliminar', {}), ('/usuarios/1/rol', {'id_rol': '2'}),
        )
        for perfil in ('cajero1', 'bodeguero'):
            self._login(perfil)
            token = self._csrf()
            antes = self._estado_ventas()
            for ruta, datos in casos:
                with self.subTest(perfil=perfil, ruta=ruta, datos=datos):
                    self._denegado(self._form(ruta, datos, token=token))
            self.assertEqual(self._estado_ventas(), antes)
            self.assertEqual(self._valor('SELECT costo FROM producto WHERE id_producto=1'), Decimal('9'))
            self.assertEqual(self._valor('SELECT precio_venta FROM producto WHERE id_producto=1'), Decimal('15'))
            self.assertTrue(self._valor('SELECT estado FROM producto WHERE id_producto=1'))

    def test_cajero_registra_clientes_pero_no_puede_editarlos(self):
        self._login('cajero1')
        respuesta = self._form('/gestion/clientes', {
            'nombre_razon_social': 'Cliente Permitido', 'nit_dpi': '1234567-K',
            'telefono': '55551234', 'correo': 'CLIENTE@EXAMPLE.TEST',
            'direccion_fiscal': 'Zona 1',
        })
        self.assertEqual(respuesta.status_code, 302)
        cliente = self._fila("SELECT * FROM cliente WHERE nit_dpi='1234567-K'")
        self.assertEqual(cliente['correo'], 'cliente@example.test')
        self._denegado(self._form('/gestion/clientes', {
            'id': cliente['id_cliente'], 'nombre_razon_social': 'Editado',
        }))
        self.assertEqual(self._valor('SELECT nombre_razon_social FROM cliente WHERE id_cliente=:id',
                                     id=cliente['id_cliente']), 'Cliente Permitido')
        self._login('bodeguero')
        self._denegado(self._form('/gestion/clientes', {'nombre_razon_social': 'Prohibido'}))

    def test_cajero_conserva_validaciones_al_crear_cliente(self):
        self._login('cajero1')
        antes = self._valor('SELECT COUNT(*) FROM cliente')
        for datos in ({'nombre_razon_social': ''},
                      {'nombre_razon_social': 'Cliente Válido', 'correo': 'correo-invalido'},
                      {'nombre_razon_social': 'Cliente Válido', 'telefono': 'abcdef'}):
            with self.subTest(datos=datos):
                self.assertEqual(self._form('/gestion/clientes', datos).status_code, 302)
                self.assertEqual(self._valor('SELECT COUNT(*) FROM cliente'), antes)

    def test_administrador_modifica_precios_costos_y_aplica_descuento(self):
        self._login('admin')
        self.assertEqual(self._form('/productos/1/editar', {
            'id_categoria': '1', 'codigo': 'PG001', 'nombre': 'Cuaderno Uno',
            'marca': 'Norma', 'unidad': 'Unidad', 'costo': '10.50',
            'precio_venta': '20', 'stock_minimo': '5', 'estado': 'on',
        }).status_code, 302)
        producto = self._fila('SELECT costo,precio_venta FROM producto WHERE id_producto=1')
        self.assertEqual(producto['costo'], Decimal('10.50'))
        self.assertEqual(producto['precio_venta'], Decimal('20'))
        self.assertEqual(self._form('/caja', {
            'accion': 'abrir', 'fondo_inicial': '100',
        }).status_code, 302)
        respuesta = self._cobrar(items=[{'id': 1, 'cantidad': 1, 'descuento': '2.50'}])
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.get_json()['total'], '17.50')
        detalle = self._fila('SELECT * FROM detalle_venta')
        self.assertEqual(detalle['descuento'], Decimal('2.50'))
        self.assertEqual(detalle['costo_unitario_historico'], Decimal('10.50'))
        self.assertEqual(self._valor('SELECT id_usuario FROM venta'), 1)

    def test_ventas_atribuidas_al_usuario_del_servidor(self):
        self._login('cajero1')
        respuesta = self._cobrar(id_usuario=3, usuario_id=3, rol='Administrador')
        self.assertEqual(respuesta.status_code, 200)
        venta = self._fila('SELECT * FROM venta WHERE id_venta=:id', id=respuesta.get_json()['id_venta'])
        self.assertEqual(venta['id_usuario'], 2)
        self.assertEqual(venta['id_turno'], 1)

    def test_cajero_no_puede_inyectar_descuentos_ni_precios_desde_json(self):
        self._login('cajero1')
        antes = self._estado_ventas()
        for campos in ({'descuento': '1'}, {'precio': '0.01'},
                       {'precio_unitario': '0.01'}, {'costo': '0.01'},
                       {'costo_unitario_historico': '0.01'}):
            with self.subTest(campos=campos):
                self._denegado(self._cobrar(items=[{'id': 1, 'cantidad': 1, **campos}]))
                self.assertEqual(self._estado_ventas(), antes)
        self.assertEqual(self._stock(), 10)

    def test_bodeguero_no_cobra_ni_administra_caja(self):
        self._login('bodeguero')
        self._denegado(self._cobrar())
        self._denegado(self._form('/caja', {'accion': 'abrir', 'fondo_inicial': '100'}))
        self.assertEqual(self._valor('SELECT COUNT(*) FROM venta'), 0)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM turno_caja WHERE id_usuario=4'), 0)

    def test_cajero_solo_consulta_e_imprime_sus_ventas_incluso_con_ids_ajenos(self):
        propia = self._venta('cajero1', producto=1)
        ajena = self._venta('cajero2', producto=2, cantidad=2)
        self._login('cajero1')
        listado = self.client.get('/ventas?id_usuario=3&rol=Administrador').get_data(as_text=True)
        self.assertIn(propia['numero'], listado)
        self.assertNotIn(ajena['numero'], listado)
        self.assertEqual(self.client.get(f"/ventas/{propia['id_venta']}/ticket").status_code, 200)
        for sufijo in ('', '?id_usuario=2&rol=Administrador', '?imprimir=1'):
            self._denegado(self.client.get(f"/ventas/{ajena['id_venta']}/ticket{sufijo}"))
        self._login('bodeguero')
        self._denegado(self.client.get(f"/ventas/{propia['id_venta']}/ticket"))
        self._login('admin')
        for venta in (propia, ajena):
            self.assertEqual(self.client.get(f"/ventas/{venta['id_venta']}/ticket").status_code, 200)

    def test_reportes_y_ambas_exportaciones_solo_incluyen_ventas_propias(self):
        propia = self._venta('cajero1', producto=1)
        ajena = self._venta('cajero2', producto=2, cantidad=2)
        self._login('cajero1')
        html = self.client.get('/reportes?id_usuario=3').get_data(as_text=True)
        self.assertIn('Cuaderno Uno', html)
        self.assertNotIn('Cuaderno Dos', html)
        self.assertNotIn('Q 45.00', html)
        self.assertNotRegex(html.lower(), r'<th[^>]*>\s*(costo|ganancia)')
        for ruta in ('/reportes/excel', '/reportes/ventas/excel'):
            respuesta = self.client.get(ruta + '?id_usuario=3&rol=Administrador')
            self.assertEqual(respuesta.status_code, 200)
            libro = load_workbook(BytesIO(respuesta.data), read_only=True, data_only=False)
            filas = list(libro.active.iter_rows(values_only=True))
            libro.close()
            self.assertEqual(len(filas), 2)
            self.assertEqual(filas[1][0], propia['numero'])
            self.assertNotIn(ajena['numero'], str(filas))
            self.assertFalse({'Costo', 'Ganancia'} & set(filas[0]))
        self._login('admin')
        respuesta = self.client.get('/reportes/excel')
        libro = load_workbook(BytesIO(respuesta.data), read_only=True)
        tickets = [fila[0] for fila in libro.active.iter_rows(min_row=2, values_only=True)]
        libro.close()
        self.assertCountEqual(tickets, [propia['numero'], ajena['numero']])

    def test_costos_y_ganancias_no_se_exponen_en_las_paginas_permitidas(self):
        self._venta('cajero1')
        for perfil, rutas in (
            ('cajero1', ('/productos', '/pos', '/ventas/1/ticket', '/reportes')),
            ('bodeguero', ('/productos', '/inventario', '/gestion/proveedores')),
        ):
            self._login(perfil)
            for ruta in rutas:
                with self.subTest(perfil=perfil, ruta=ruta):
                    respuesta = self.client.get(ruta)
                    self.assertEqual(respuesta.status_code, 200)
                    html = respuesta.get_data(as_text=True)
                    self.assertNotIn('Q 9.00', html)
                    self.assertNotIn('costo_unitario_historico', html)
                    self.assertNotRegex(html.lower(), r'<th[^>]*>\s*(costo|ganancia)')
        self._login('admin')
        self.assertIn('Q 9.00', self.client.get('/productos').get_data(as_text=True))

    def test_menus_y_acciones_se_ajustan_al_rol(self):
        self._login('cajero1')
        html = self.client.get('/productos').get_data(as_text=True)
        for ruta in ('/inventario', '/usuarios', '/configuracion', '/compras',
                     '/gestion/categorias', '/gestion/proveedores', '/productos/nuevo',
                     '/productos/1/editar', '/productos/1/estado', '/productos/1/eliminar'):
            self.assertNotIn(f'href="{ruta}"', html)
            self.assertNotIn(f'action="{ruta}"', html)
        self._login('bodeguero')
        html = self.client.get('/productos').get_data(as_text=True)
        for ruta in ('/pos', '/ventas', '/usuarios', '/reportes', '/caja',
                     '/gestion/clientes', '/productos/nuevo'):
            self.assertNotIn(f'href="{ruta}"', html)
        self.assertIn('href="/inventario"', html)
        self.assertIn('href="/gestion/proveedores"', html)

    def test_cambio_de_rol_se_refleja_en_la_siguiente_solicitud(self):
        # La regla de caja vigente exige cerrar el turno antes de pasar al almacén.
        self._sql("UPDATE turno_caja SET estado='CERRADO',cierre=NOW() WHERE id_usuario=2")
        cliente_admin = self._login('admin')
        cliente_cajero = self._login('cajero1')
        self.assertEqual(cliente_cajero.get('/pos').status_code, 200)
        respuesta = self._form('/usuarios/2/rol', {'id_rol': '3'}, cliente=cliente_admin)
        self.assertEqual(respuesta.status_code, 302)
        self._denegado(cliente_cajero.get('/pos'))
        self.assertEqual(cliente_cajero.get('/inventario').status_code, 200)
        respuesta = self._form('/usuarios/2/rol', {'id_rol': '2'}, cliente=cliente_admin)
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(cliente_cajero.get('/pos').status_code, 200)
        self._denegado(cliente_cajero.get('/inventario'))

    def test_usuario_desactivado_pierde_acceso_en_su_siguiente_solicitud(self):
        self._sql("UPDATE turno_caja SET estado='CERRADO',cierre=NOW() WHERE id_usuario=2")
        administrador = self._login('admin')
        cajero = self._login('cajero1')
        token = self._csrf(cajero)
        respuesta = self._form('/usuarios/2/estado', {'estado': 'desactivar'}, cliente=administrador)
        self.assertEqual(respuesta.status_code, 302)
        self.assertFalse(self._valor('SELECT estado FROM usuario WHERE id_usuario=2'))
        self.assertEqual(cajero.get('/productos').status_code, 302)
        respuesta = cajero.post('/pos/cobrar', json=self._payload(), headers={'X-CSRFToken': token})
        self.assertIn(respuesta.status_code, (302, 401))
        self.assertEqual(self._valor('SELECT COUNT(*) FROM venta'), 0)
        nuevo = self.app.app.test_client()
        respuesta = nuevo.post('/login', data={
            'correo': 'cajero1@example.test', 'password': 'PruebaPostgres2026!',
            'csrf_token': self._csrf(nuevo, '/login'),
        })
        self.assertEqual(respuesta.status_code, 200)

    def test_rol_del_navegador_no_concede_permisos(self):
        self._login('cajero1')
        with self.client.session_transaction() as sesion:
            sesion['rol'] = 'Administrador'
            sesion['id_rol'] = 1
        self._denegado(self.client.get('/usuarios?rol=Administrador&id_rol=1'))
        self._denegado(self._form('/usuarios', {
            'nombre': 'Administrador Falso', 'correo': 'falso@example.test',
            'password': 'PruebaPostgres2026!', 'id_rol': '1', 'rol': 'Administrador',
        }))
        self.assertEqual(self._valor('SELECT COUNT(*) FROM usuario'), 4)

    def test_ultimo_admin_no_puede_perder_el_rol_y_rol_desconocido_se_rechaza(self):
        self._login('admin')
        for rol in ('2', '3', '999999'):
            respuesta = self._form('/usuarios/1/rol', {'id_rol': rol})
            self.assertEqual(respuesta.status_code, 302)
            self.assertEqual(self._valor('SELECT id_rol FROM usuario WHERE id_usuario=1'), 1)
        desconocido = self._crear("INSERT INTO rol(nombre) VALUES('Rol sin permisos') RETURNING id_rol")
        self.assertEqual(self._form('/usuarios/2/rol', {'id_rol': str(desconocido)}).status_code, 302)
        self.assertEqual(self._valor('SELECT id_rol FROM usuario WHERE id_usuario=2'), 2)
        self.assertEqual(self.client.get('/usuarios').status_code, 200)

    def test_admin_no_elimina_ni_desactiva_su_cuenta_aun_con_otro_admin(self):
        self._otro_admin()
        self._login('admin')
        for ruta in ('/usuarios/1/eliminar', '/usuarios/1/estado'):
            respuesta = self._form(ruta, {'estado': 'desactivar'})
            self.assertEqual(respuesta.status_code, 302)
            self.assertTrue(self._valor('SELECT estado FROM usuario WHERE id_usuario=1'))
        self.assertEqual(self._valor('SELECT COUNT(*) FROM usuario WHERE id_usuario=1'), 1)

    def test_dos_cambios_concurrentes_no_dejan_al_sistema_sin_administrador(self):
        otro = self._otro_admin()
        primero = self._login('admin')
        segundo = self._login('otroadmin')
        token1, token2 = self._csrf(primero), self._csrf(segundo)
        with ThreadPoolExecutor(max_workers=2) as pool:
            operaciones = (
                pool.submit(self._form, '/usuarios/1/rol', {'id_rol': '2'},
                            cliente=primero, token=token1),
                pool.submit(self._form, f'/usuarios/{otro}/rol', {'id_rol': '2'},
                            cliente=segundo, token=token2),
            )
            for operacion in operaciones:
                self.assertIn(operacion.result(timeout=20).status_code, (302, 403))
        self.assertEqual(self._valor('''SELECT COUNT(*) FROM usuario u JOIN rol r USING(id_rol)
            WHERE u.estado=TRUE AND r.nombre='Administrador' '''), 1)

    def test_admin_elimina_registros_libres_y_conserva_registros_con_historial(self):
        categoria = self._crear("INSERT INTO categoria(nombre) VALUES('Categoría libre') RETURNING id_categoria")
        cliente = self._crear("INSERT INTO cliente(nombre_razon_social) VALUES('Cliente Libre') RETURNING id_cliente")
        proveedor = self._crear("INSERT INTO proveedor(razon_social) VALUES('Proveedor Libre') RETURNING id_proveedor")
        usuario = self._crear('''INSERT INTO usuario(id_rol,nombre,correo,hash_contrasena)
            VALUES(2,'Usuario Libre','libre@example.test',:hash) RETURNING id_usuario''', hash=HASH_PRUEBA)
        producto = self._crear('''INSERT INTO producto(id_categoria,codigo,nombre,unidad,precio_venta)
            VALUES(1,'LIBRE','Producto libre','unidad',10) RETURNING id_producto''')
        self._login('admin')
        casos = (
            (f'/gestion/categorias/{categoria}/eliminar', 'categoria', 'id_categoria', categoria),
            (f'/gestion/clientes/{cliente}/eliminar', 'cliente', 'id_cliente', cliente),
            (f'/gestion/proveedores/{proveedor}/eliminar', 'proveedor', 'id_proveedor', proveedor),
            (f'/usuarios/{usuario}/eliminar', 'usuario', 'id_usuario', usuario),
            (f'/productos/{producto}/eliminar', 'producto', 'id_producto', producto),
        )
        for ruta, tabla, pk, rid in casos:
            with self.subTest(ruta=ruta):
                self.assertEqual(self._form(ruta).status_code, 302)
                self.assertEqual(self._valor(f'SELECT COUNT(*) FROM {tabla} WHERE {pk}=:id', id=rid), 0)
        self.assertEqual(self._form('/gestion/categorias/1/eliminar').status_code, 302)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM categoria WHERE id_categoria=1'), 1)
        self.assertEqual(self._form('/productos/1/eliminar').status_code, 302)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM producto WHERE id_producto=1'), 1)
        self.assertEqual(self._form('/usuarios/2/eliminar').status_code, 302)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM usuario WHERE id_usuario=2'), 1)
        self.assertEqual(self._stock(), 10)

    def test_venta_protege_cliente_usuario_y_producto_frente_a_eliminacion(self):
        venta = self._venta()
        self._login('admin')
        for ruta in ('/gestion/clientes/1/eliminar', '/productos/1/eliminar', '/usuarios/2/eliminar'):
            self.assertEqual(self._form(ruta).status_code, 302)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM venta WHERE id_venta=:id', id=venta['id_venta']), 1)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM detalle_venta'), 1)
        self.assertEqual(self._stock(), 9)

    def test_proveedor_con_compra_se_desactiva_y_no_se_elimina_su_historial(self):
        self._login('admin')
        self.assertEqual(self._form('/compras/nueva', {
            'id_proveedor': 1, 'id_producto': 1, 'cantidad': '2',
            'costo_unitario': '9', 'metodo_pago': 'TRANSFERENCIA',
            'numero_factura': 'HISTORIAL-PROVEEDOR',
        }).status_code, 302)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM compra'), 1)
        self.assertEqual(self._form('/gestion/proveedores/1/eliminar').status_code, 302)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM proveedor WHERE id_proveedor=1'), 1)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM compra'), 1)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM detalle_compra'), 1)
        self.assertEqual(self._form('/gestion/proveedores/1/estado', {'estado': 'desactivar'}).status_code, 302)
        self.assertFalse(self._valor('SELECT estado FROM proveedor WHERE id_proveedor=1'))
        self.assertEqual(self._stock(), 12)
        self._login('bodeguero')
        self.assertEqual(self.client.get('/gestion/proveedores').status_code, 200)

    def test_admin_desactiva_registros_relacionados_sin_borrar_historial(self):
        comprador = self._crear("INSERT INTO cliente(nombre_razon_social) VALUES('Cliente Histórico') RETURNING id_cliente")
        venta = self._venta(cliente_id=comprador)
        self._sql("UPDATE turno_caja SET estado='CERRADO',cierre=NOW() WHERE id_usuario=2")
        self._login('admin')
        for ruta, tabla, pk, rid in (
            ('/gestion/categorias/1/estado', 'categoria', 'id_categoria', 1),
            (f'/gestion/clientes/{comprador}/estado', 'cliente', 'id_cliente', comprador),
            ('/usuarios/2/estado', 'usuario', 'id_usuario', 2),
            ('/productos/1/estado', 'producto', 'id_producto', 1),
        ):
            with self.subTest(ruta=ruta):
                self.assertEqual(self._form(ruta, {'estado': 'desactivar'}).status_code, 302)
                self.assertFalse(self._valor(f'SELECT estado FROM {tabla} WHERE {pk}=:id', id=rid))
        self.assertEqual(self._valor('SELECT COUNT(*) FROM venta WHERE id_venta=:id', id=venta['id_venta']), 1)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM detalle_venta'), 1)
        self.assertEqual(self._stock(), 9)

    def test_anulacion_solo_admin_conserva_historial_y_no_restituye_stock_dos_veces(self):
        venta = self._venta()
        ruta = f"/ventas/{venta['id_venta']}/anular"
        self._denegado(self._form(ruta, {'motivo': 'Anulación no permitida'}))
        self._login('bodeguero')
        self._denegado(self._form(ruta, {'motivo': 'Anulación no permitida'}))
        self.assertEqual(self._stock(), 9)
        self._login('admin')
        self.assertEqual(self._form('/caja', {
            'accion': 'abrir', 'fondo_inicial': '100',
        }).status_code, 302)
        self.assertEqual(self._form(ruta, {'motivo': 'Error de captura'}).status_code, 302)
        self.assertEqual(self._valor('SELECT estado FROM venta WHERE id_venta=:id', id=venta['id_venta']), 'ANULADA')
        self.assertEqual(self._stock(), 10)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM detalle_venta'), 1)
        self.assertEqual(self._form(ruta, {'motivo': 'Reintento'}).status_code, 302)
        self.assertEqual(self._stock(), 10)

    def test_bodeguero_registra_movimiento_con_responsable_servidor_y_no_cambia_costos(self):
        self._login('bodeguero')
        respuesta = self._form('/inventario/movimiento', {
            'id_producto': 1, 'tipo': 'ENTRADA', 'cantidad': '3',
            'motivo': 'Recepción de mercadería', 'id_usuario': 1, 'costo_unitario': '0.01',
        })
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(self._stock(), 13)
        movimiento = self._fila('''SELECT * FROM movimiento_inventario
            WHERE motivo='Recepción de mercadería' ''')
        self.assertEqual(movimiento['id_usuario'], 4)
        self.assertEqual(movimiento['id_producto'], 1)
        self.assertEqual(movimiento['cantidad'], 3)
        self.assertIsNotNone(movimiento['fecha_hora'])
        self.assertEqual(self._valor('SELECT costo FROM producto WHERE id_producto=1'), Decimal('9'))
        self.assertEqual(self._form('/inventario/movimiento', {
            'id_producto': 1, 'tipo': 'SALIDA', 'cantidad': '4', 'motivo': 'Despacho de mercadería',
        }).status_code, 302)
        self.assertEqual(self._stock(), 9)

    def test_salida_no_deja_stock_negativo_y_cajero_no_registra_movimientos(self):
        self._login('bodeguero')
        for cantidad in ('11', '-1', '1.5'):
            self.assertEqual(self._form('/inventario/movimiento', {
                'id_producto': 1, 'tipo': 'SALIDA', 'cantidad': cantidad, 'motivo': 'Salida inválida',
            }).status_code, 302)
            self.assertEqual(self._stock(), 10)
        self._login('cajero1')
        self._denegado(self._form('/inventario/movimiento', {
            'id_producto': 1, 'tipo': 'ENTRADA', 'cantidad': '3', 'motivo': 'Entrada prohibida',
        }))
        self.assertEqual(self._stock(), 10)

    def test_salidas_concurrentes_bloquean_producto_y_no_sobregiran_existencias(self):
        primero, segundo = self._login('bodeguero'), self._login('bodeguero')
        token1, token2 = self._csrf(primero), self._csrf(segundo)
        datos = {'id_producto': 1, 'tipo': 'SALIDA', 'cantidad': '7'}
        with ThreadPoolExecutor(max_workers=2) as pool:
            operaciones = (
                pool.submit(self._form, '/inventario/movimiento', {
                    **datos, 'motivo': 'Primera salida simultánea',
                }, cliente=primero, token=token1),
                pool.submit(self._form, '/inventario/movimiento', {
                    **datos, 'motivo': 'Segunda salida simultánea',
                }, cliente=segundo, token=token2),
            )
            for operacion in operaciones:
                self.assertEqual(operacion.result(timeout=20).status_code, 302)
        self.assertEqual(self._stock(), 3)
        self.assertEqual(self._valor("SELECT COUNT(*) FROM movimiento_inventario WHERE tipo='SALIDA'"), 1)

    def _solicitar_ajuste(self, cantidad=7):
        self._login('bodeguero')
        respuesta = self._form('/inventario/conteos', {
            'id_producto': 1, 'cantidad_contada': cantidad,
            'motivo': 'Diferencia en conteo físico', 'id_usuario': 1,
        })
        self.assertEqual(respuesta.status_code, 302)
        return self._fila('SELECT * FROM solicitud_ajuste ORDER BY id_solicitud DESC LIMIT 1')

    def test_conteo_solicita_ajuste_pendiente_sin_cambiar_stock(self):
        solicitud = self._solicitar_ajuste()
        self.assertEqual(solicitud['estado'], 'PENDIENTE')
        self.assertEqual(solicitud['id_solicitante'], 4)
        self.assertEqual(solicitud['existencia_sistema'], 10)
        self.assertEqual(solicitud['cantidad_contada'], 7)
        self.assertEqual(solicitud['diferencia'], -3)
        self.assertEqual(self._stock(), 10)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM movimiento_inventario'), 2)
        for accion in ('aprobar', 'rechazar'):
            self._denegado(self._form(f"/inventario/ajustes/{solicitud['id_solicitud']}/{accion}"))
        self.assertEqual(self._stock(), 10)
        self._login('cajero1')
        self._denegado(self._form('/inventario/conteos', {
            'id_producto': 1, 'cantidad_contada': '7', 'motivo': 'Conteo prohibido',
        }))
        self._denegado(self._form(f"/inventario/ajustes/{solicitud['id_solicitud']}/aprobar"))

    def test_solo_admin_aplica_ajuste_una_vez(self):
        solicitud = self._solicitar_ajuste()
        self._login('admin')
        ruta = f"/inventario/ajustes/{solicitud['id_solicitud']}/aprobar"
        self.assertEqual(self._form(ruta).status_code, 302)
        self.assertEqual(self._stock(), 7)
        self.assertEqual(self._valor('SELECT estado FROM solicitud_ajuste WHERE id_solicitud=:id',
                                     id=solicitud['id_solicitud']), 'APROBADA')
        self.assertEqual(self._valor('SELECT id_aprobador FROM solicitud_ajuste WHERE id_solicitud=:id',
                                     id=solicitud['id_solicitud']), 1)
        movimientos = self._valor('SELECT COUNT(*) FROM movimiento_inventario')
        self.assertEqual(self._form(ruta).status_code, 302)
        self.assertEqual(self._stock(), 7)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM movimiento_inventario'), movimientos)

    def test_aprobaciones_concurrentes_del_mismo_ajuste_lo_aplican_una_sola_vez(self):
        solicitud = self._solicitar_ajuste()
        primero, segundo = self._login('admin'), self._login('admin')
        token1, token2 = self._csrf(primero), self._csrf(segundo)
        ruta = f"/inventario/ajustes/{solicitud['id_solicitud']}/aprobar"
        with ThreadPoolExecutor(max_workers=2) as pool:
            operaciones = (
                pool.submit(self._form, ruta, cliente=primero, token=token1),
                pool.submit(self._form, ruta, cliente=segundo, token=token2),
            )
            for operacion in operaciones:
                self.assertEqual(operacion.result(timeout=20).status_code, 302)
        self.assertEqual(self._stock(), 7)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM movimiento_inventario'), 3)

    def test_ajuste_rechazado_conserva_historial_y_no_puede_aplicarse(self):
        solicitud = self._solicitar_ajuste()
        self._login('admin')
        base = f"/inventario/ajustes/{solicitud['id_solicitud']}"
        self.assertEqual(self._form(base + '/rechazar', {'motivo': 'Conteo incorrecto'}).status_code, 302)
        self.assertEqual(self._valor('SELECT estado FROM solicitud_ajuste WHERE id_solicitud=:id',
                                     id=solicitud['id_solicitud']), 'RECHAZADA')
        self.assertEqual(self._form(base + '/aprobar').status_code, 302)
        self.assertEqual(self._stock(), 10)
        self.assertEqual(self._valor('SELECT COUNT(*) FROM movimiento_inventario'), 2)

    def test_ajuste_pendiente_no_se_aplica_si_stock_cambio_despues_del_conteo(self):
        solicitud = self._solicitar_ajuste()
        self._venta('cajero1')
        self.assertEqual(self._stock(), 9)
        self._login('admin')
        self.assertEqual(self._form(
            f"/inventario/ajustes/{solicitud['id_solicitud']}/aprobar").status_code, 302)
        self.assertEqual(self._stock(), 9)
        self.assertEqual(self._valor('SELECT estado FROM solicitud_ajuste WHERE id_solicitud=:id',
                                     id=solicitud['id_solicitud']), 'PENDIENTE')

    def test_conteo_invalido_no_crea_solicitud_y_cero_se_aplica_sin_stock_negativo(self):
        self._login('bodeguero')
        for cantidad in ('-1', '1.5', 'NaN'):
            with self.subTest(cantidad=cantidad):
                self.assertEqual(self._form('/inventario/conteos', {
                    'id_producto': 1, 'cantidad_contada': cantidad, 'motivo': 'Conteo inválido',
                }).status_code, 302)
                self.assertEqual(self._valor('SELECT COUNT(*) FROM solicitud_ajuste'), 0)
                self.assertEqual(self._stock(), 10)
        solicitud = self._solicitar_ajuste(cantidad=0)
        self._login('admin')
        self.assertEqual(self._form(
            f"/inventario/ajustes/{solicitud['id_solicitud']}/aprobar").status_code, 302)
        self.assertEqual(self._stock(), 0)

    def test_exportacion_inventario_no_revela_costos_a_bodeguero(self):
        for perfil in ('admin', 'bodeguero'):
            self._login(perfil)
            respuesta = self.client.get('/inventario/excel')
            self.assertEqual(respuesta.status_code, 200)
            libro = load_workbook(BytesIO(respuesta.data), read_only=True)
            encabezados = [str(celda).lower() for celda in next(libro.active.iter_rows(values_only=True))]
            libro.close()
            if perfil == 'bodeguero':
                self.assertFalse(any('costo' in campo or 'ganancia' in campo for campo in encabezados))
            else:
                self.assertTrue(any('costo' in campo for campo in encabezados))


