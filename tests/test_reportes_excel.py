"""Pruebas aisladas: no importan app ni necesitan PostgreSQL o configuración local."""

import unittest
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from unittest.mock import patch

from openpyxl import load_workbook

from reportes_excel import generar_excel_ventas, resolver_rango_fechas, validar_rango_fechas


class RangosFechasTests(unittest.TestCase):
    def test_acepta_un_dia_y_366_dias_incluidos(self):
        self.assertEqual(validar_rango_fechas("2024-02-29", "2024-02-29"), (date(2024, 2, 29),) * 2)
        self.assertEqual(validar_rango_fechas("2024-01-01", "2024-12-31"), (date(2024, 1, 1), date(2024, 12, 31)))

    def test_rechaza_fechas_inexistentes_rangos_invertidos_y_excesivos(self):
        for inicio, fin in (
            ("2025-02-29", "2025-03-01"), ("2026-02-02", "2026-02-01"),
            ("2024-01-01", "2025-01-01"), ("20260101", "20260102"),
            ("2026-W01-1", "2026-W01-2"), ("", "2026-01-01"),
            (None, None), ("9999-12-31", "9999-12-31"),
        ):
            with self.subTest(inicio=inicio, fin=fin), self.assertRaises(ValueError):
                validar_rango_fechas(inicio, fin)

    def test_default_usa_el_hoy_recibido_y_cruza_el_ano(self):
        self.assertEqual(resolver_rango_fechas({}, hoy=date(2026, 1, 15)), (date(2025, 12, 17), date(2026, 1, 15)))

    def test_ambos_pares_de_parametros_producen_el_mismo_rango(self):
        esperado = (date(2026, 10, 1), date(2026, 10, 17))
        for parametros in (
            {"desde": "2026-10-01", "hasta": "2026-10-17"},
            {"inicio": "2026-10-01", "fin": "2026-10-17"},
            {"desde": "2026-10-01", "inicio": "2026-10-01", "hasta": "2026-10-17", "fin": "2026-10-17"},
        ):
            with self.subTest(parametros=parametros):
                self.assertEqual(resolver_rango_fechas(parametros, hoy=date(2020, 1, 1)), esperado)

    def test_rechaza_parametros_contradictorios_y_rangos_incompletos(self):
        for parametros in (
            {"desde": "2026-10-01", "inicio": "2026-09-01", "hasta": "2026-10-17"},
            {"inicio": "2026-10-01"}, {"hasta": "2026-10-17"},
        ):
            with self.subTest(parametros=parametros), self.assertRaises(ValueError):
                resolver_rango_fechas(parametros, hoy=date(2026, 10, 17))


class ExcelVentasTests(unittest.TestCase):
    @staticmethod
    def venta(**cambios):
        venta = {
            "numero_ticket": "000001", "fecha_hora": datetime(2026, 10, 17, 8, 30),
            "cliente": "Cliente de prueba", "cajero": "Ana", "metodo_pago": "EFECTIVO",
            "total": Decimal("123.45"), "estado": "COMPLETADA",
        }
        venta.update(cambios)
        return venta

    def test_vacio_es_excel_valido_y_tiene_formato(self):
        archivo = generar_excel_ventas([])
        self.assertEqual(archivo.tell(), 0)
        libro = load_workbook(archivo)
        hoja = libro.active
        self.assertEqual(hoja.max_row, 1)
        self.assertEqual(hoja["A1"].value, "Ticket")
        self.assertTrue(hoja["A1"].font.bold)
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(hoja.auto_filter.ref, "A1:G1")

    def test_fecha_con_zona_se_convierte_a_guatemala_y_cambia_el_dia(self):
        venta = self.venta(fecha_hora=datetime(2026, 10, 17, 2, 30, tzinfo=timezone.utc))
        hoja = load_workbook(generar_excel_ventas([venta])).active
        self.assertEqual(hoja["B2"].value, datetime(2026, 10, 16, 20, 30))
        self.assertIsNone(hoja["B2"].value.tzinfo)
        self.assertEqual(hoja["F2"].value, 123.45)
        self.assertEqual(hoja["F2"].number_format, '"Q" #,##0.00')
        self.assertEqual(hoja["A2"].value, "000001")

    def test_fecha_local_se_preserva_y_zona_explicita_es_respetada(self):
        local = self.venta()
        hoja = load_workbook(generar_excel_ventas([local])).active
        self.assertEqual(hoja["B2"].value, local["fecha_hora"])
        aware = self.venta(fecha_hora=datetime(2026, 10, 17, 0, 0, tzinfo=timezone.utc))
        hoja = load_workbook(generar_excel_ventas([aware], zona_horaria=timezone(timedelta(hours=2)))).active
        self.assertEqual(hoja["B2"].value, datetime(2026, 10, 17, 2, 0))

    def test_texto_de_usuario_no_se_convierte_en_formula_o_error(self):
        for texto in ('=HYPERLINK("https://example.invalid", "Abrir")', "+SUM(1,2)", "-1+2", "@SUM(A1)", "#N/A"):
            with self.subTest(texto=texto):
                venta = self.venta(numero_ticket=texto, cliente=texto, cajero=texto, metodo_pago=texto, estado=texto)
                hoja = load_workbook(generar_excel_ventas([venta]), data_only=False).active
                for columna in (1, 3, 4, 5, 7):
                    celda = hoja.cell(2, columna)
                    self.assertEqual(celda.value, texto)
                    self.assertEqual(celda.data_type, "s")

    def test_no_trunca_despues_de_5000_ventas_y_acepta_generadores(self):
        ventas = (self.venta(numero_ticket=str(indice)) for indice in range(5001))
        hoja = load_workbook(generar_excel_ventas(ventas), read_only=True).active
        self.assertEqual(hoja.max_row, 5002)
        self.assertEqual(hoja["A5002"].value, "5000")

    def test_limite_de_excel_es_error_explicito(self):
        with patch("reportes_excel.MAX_FILAS_EXCEL", 3):
            with self.assertRaisesRegex(ValueError, "capacidad de Excel"):
                generar_excel_ventas(self.venta() for _ in range(3))

    def test_datos_invalidos_no_se_silencian(self):
        for cambios in (
            {"total": "NaN"}, {"total": "Infinity"}, {"total": "incorrecto"},
            {"fecha_hora": "2026-10-17"}, {"cliente": "\x00"}, {"cliente": "a" * 32768},
        ):
            with self.subTest(cambios=cambios), self.assertRaises(ValueError):
                generar_excel_ventas([self.venta(**cambios)])


if __name__ == "__main__":
    unittest.main()
