"""Rangos y exportación de ventas, sin dependencias de Flask ni de la base de datos."""

import re
from datetime import date, datetime, timedelta, tzinfo
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Iterable, Mapping
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


MAX_DIAS_REPORTE = 366
MAX_FILAS_EXCEL = 1_048_576  # Incluye el encabezado.
_FECHA_ISO = re.compile(r"\d{4}-\d{2}-\d{2}\Z")
_CARACTERES_INVALIDOS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_ENCABEZADOS = (
    "Ticket", "Fecha", "Cliente", "Cajero", "Método de pago", "Total Q", "Estado"
)


def validar_rango_fechas(inicio: str, fin: str) -> tuple[date, date]:
    """Valida fechas YYYY-MM-DD y hasta 366 días incluidos, sin consultar el reloj."""
    mensaje = "Selecciona un rango de fechas válido (AAAA-MM-DD)."
    if not isinstance(inicio, str) or not isinstance(fin, str):
        raise ValueError(mensaje)
    if not _FECHA_ISO.fullmatch(inicio) or not _FECHA_ISO.fullmatch(fin):
        raise ValueError(mensaje)
    try:
        desde = date.fromisoformat(inicio)
        hasta = date.fromisoformat(fin)
    except ValueError as error:
        raise ValueError(mensaje) from error
    if hasta < desde:
        raise ValueError("La fecha final no puede ser anterior a la inicial.")
    if (hasta - desde).days + 1 > MAX_DIAS_REPORTE:
        raise ValueError(f"El rango no puede superar {MAX_DIAS_REPORTE} días.")
    # El SQL utiliza una cota final exclusiva, que debe poder representarse.
    if hasta == date.max:
        raise ValueError("La fecha final está fuera del rango permitido.")
    return desde, hasta


def resolver_rango_fechas(
    parametros: Mapping[str, str], *, hoy: date
) -> tuple[date, date]:
    """Acepta desde/hasta e inicio/fin; por defecto usa los últimos 30 días.

    Los alias pueden coexistir si coinciden. Un rango parcialmente indicado o
    valores contradictorios se rechazan para no exportar otro período por error.
    El llamador proporciona la fecha local para no depender del reloj del servidor.
    """
    valores = {}
    for principal, alias in (("desde", "inicio"), ("hasta", "fin")):
        primero = (parametros.get(principal) or "").strip()
        segundo = (parametros.get(alias) or "").strip()
        if primero and segundo and primero != segundo:
            raise ValueError("Los parámetros de fechas indican rangos diferentes.")
        valores[principal] = primero or segundo

    if not valores["desde"] and not valores["hasta"]:
        valores["desde"] = (hoy - timedelta(days=29)).isoformat()
        valores["hasta"] = hoy.isoformat()
    return validar_rango_fechas(valores["desde"], valores["hasta"])


def _texto_excel(celda, valor) -> None:
    """Mantiene texto literal, incluso si empieza con =, +, - o @."""
    texto = "" if valor is None else str(valor)
    if len(texto) > 32_767 or _CARACTERES_INVALIDOS.search(texto):
        raise ValueError("Una venta contiene texto incompatible con el formato Excel.")
    celda.value = texto
    # openpyxl infiere fórmulas a partir de '=' y errores como '#N/A'. La
    # asignación explícita después del valor preserva los datos como texto.
    celda.data_type = "s"


def _fecha_excel(valor: datetime, zona_horaria: tzinfo) -> datetime:
    if not isinstance(valor, datetime):
        raise ValueError("Una venta contiene una fecha y hora inválida.")
    if valor.tzinfo is not None and valor.utcoffset() is not None:
        return valor.astimezone(zona_horaria).replace(tzinfo=None)
    # Los timestamp sin zona del esquema representan la hora local del negocio.
    return valor


def generar_excel_ventas(
    ventas: Iterable[Mapping], *, zona_horaria: tzinfo | None = None
) -> BytesIO:
    """Exporta todas las filas recibidas y devuelve un archivo listo para descargar.

    Cada fila necesita numero_ticket, fecha_hora, cliente, cajero, metodo_pago,
    total y estado. La consulta decide qué estados exportar. No hay límite de
    5000 ventas; si se supera la capacidad de una hoja se informa explícitamente.
    Las fechas con zona se convierten a Guatemala antes de guardarlas sin zona,
    porque Excel no admite datetime con información de zona horaria.
    """
    if zona_horaria is None:
        zona_horaria = ZoneInfo("America/Guatemala")
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas PUNTO COM"
    hoja.append(_ENCABEZADOS)
    for celda in hoja[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor="1D4ED8")
        celda.alignment = Alignment(vertical="center")
    hoja.row_dimensions[1].height = 24

    for numero_fila, venta in enumerate(ventas, start=2):
        if numero_fila > MAX_FILAS_EXCEL:
            raise ValueError("El reporte excede la capacidad de Excel. Selecciona un período menor.")
        fecha = _fecha_excel(venta["fecha_hora"], zona_horaria)
        try:
            total = Decimal(str(venta["total"]))
        except (InvalidOperation, TypeError, ValueError) as error:
            raise ValueError("Una venta contiene un importe inválido.") from error
        if not total.is_finite():
            raise ValueError("Una venta contiene un importe inválido.")
        hoja.cell(numero_fila, 2, fecha).number_format = "DD/MM/YYYY HH:MM"
        hoja.cell(numero_fila, 6, total).number_format = '"Q" #,##0.00'
        for columna, clave in (
            (1, "numero_ticket"), (3, "cliente"), (4, "cajero"),
            (5, "metodo_pago"), (7, "estado"),
        ):
            _texto_excel(hoja.cell(numero_fila, columna), venta[clave])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for columna, ancho in {"A": 22, "B": 23, "C": 35, "D": 28, "E": 22, "F": 18, "G": 20}.items():
        hoja.column_dimensions[columna].width = ancho
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.print_title_rows = "1:1"

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return archivo
