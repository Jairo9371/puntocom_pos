"""Carga verificable de 30 operaciones PRUEBA, sin crear usuarios.

Por defecto revierte el lote completo. --aplicar lo conserva únicamente cuando
las treinta operaciones y sus comprobaciones terminan correctamente. Las rutas
HTTP, los permisos, los validadores y CSRF son los mismos de la aplicación.
"""

import argparse
import json
import re
from collections import Counter
from datetime import datetime
from decimal import Decimal
from html import unescape
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import scoped_session, sessionmaker

from configuracion import TZ


ROOT = Path(__file__).resolve().parent
ROLES = ('Administrador', 'Cajero', 'Bodeguero')
CODIGOS = ('PRUEBA-ROLES-A', 'PRUEBA-ROLES-B', 'PRUEBA-ROLES-C')


class ErrorDemostracion(ValueError):
    """Error del lote; nunca se conserva una carga parcial."""


def comprobar(condicion, mensaje):
    if not condicion:
        raise ErrorDemostracion(mensaje)


class Demostracion:
    def __init__(self, modulo):
        self.modulo = modulo
        self.clientes = {}
        self.actores = {}
        self.productos = []
        self.ventas = []
        self.solicitudes = []
        self.turnos_nuevos = []
        self.reporte = {
            'fecha': datetime.now(TZ).isoformat(),
            'aplicado': False,
            'usuarios_creados': 0,
            'roles_creados': 0,
            'autenticacion': 'Sesiones de prueba asignadas en el servidor a usuarios existentes; '
                             'sin conocer ni cambiar sus contraseñas. CSRF habilitado.',
            'preparacion': [], 'operaciones': [], 'comprobaciones': [], 'registros': {},
        }

    def filas(self, sql, **datos):
        return [dict(f) for f in self.modulo.db.session.execute(text(sql), datos).mappings()]

    def valor(self, sql, **datos):
        return self.modulo.db.session.execute(text(sql), datos).scalar_one()

    def stock(self, pid):
        return int(self.valor('SELECT existencia_actual FROM vista_inventario WHERE id_producto=:id', id=pid))

    def verificar(self, nombre, condicion, **evidencia):
        comprobar(condicion, nombre)
        self.reporte['comprobaciones'].append({'nombre': nombre, 'resultado': 'CORRECTO', **evidencia})

    def operacion(self, rol, nombre, ruta, respuesta, **evidencia):
        comprobar(respuesta.status_code in (200, 302), f'Falló {nombre}: HTTP {respuesta.status_code}.')
        numero = 1 + sum(o['rol'] == rol for o in self.reporte['operaciones'])
        self.reporte['operaciones'].append({
            'rol': rol, 'numero': numero, 'nombre': nombre, 'ruta': ruta,
            'http': respuesta.status_code, 'resultado': 'CORRECTO', **evidencia,
        })

    def token(self, rol):
        respuesta = self.clientes[rol].get('/productos')
        comprobar(respuesta.status_code == 200, f'No se pudo consultar productos como {rol}.')
        html = respuesta.get_data(as_text=True)
        token = re.search(r'name="csrf_token"\s+value="([^"]+)"', html)
        if not token:
            token = re.search(r'data-csrf-token="([^"]+)"', html)
        comprobar(token is not None, 'No se encontró el token CSRF de la aplicación.')
        return unescape(token.group(1))

    def formulario(self, rol, ruta, **datos):
        return self.clientes[rol].post(ruta, data={**datos, 'csrf_token': self.token(rol)})

    def cobrar(self, rol, cliente, producto, cantidad=1, descuento='0', metodo='TRANSFERENCIA'):
        return self.clientes[rol].post('/pos/cobrar', json={
            'cliente_id': cliente, 'metodo': metodo, 'recibido': '0',
            'items': [{'id': producto, 'cantidad': cantidad, 'descuento': descuento}],
        }, headers={'X-CSRFToken': self.token(rol), 'Idempotency-Key': uuid4().hex})

    def denegado(self, nombre, respuesta):
        contenido = str(respuesta.get_json()) if respuesta.is_json else respuesta.get_data(as_text=True)
        self.verificar(nombre, respuesta.status_code == 403
                       and 'No tienes permiso para realizar esta acción' in contenido, http=respuesta.status_code)

    def preparar(self):
        comprobar(self.valor('SELECT COUNT(*) FROM producto WHERE codigo IN (:a,:b,:c)',
                             a=CODIGOS[0], b=CODIGOS[1], c=CODIGOS[2]) == 0,
                  'Ya existe la demostración PRUEBA-ROLES. No se duplicaron datos; consulta el informe anterior.')
        self.usuarios_antes = self.valor('SELECT COUNT(*) FROM usuario')
        self.roles_antes = self.valor('SELECT COUNT(*) FROM rol')
        self.caja_antes = self.valor('SELECT COUNT(*) FROM movimiento_caja')
        for rol in ROLES:
            actores = self.filas('''SELECT u.id_usuario FROM usuario u JOIN rol r USING(id_rol)
                                   WHERE u.estado=TRUE AND r.nombre=:rol ORDER BY u.id_usuario LIMIT 1''', rol=rol)
            comprobar(bool(actores), f'Se necesita un usuario activo existente con rol {rol}.')
            uid = actores[0]['id_usuario']
            self.actores[rol] = uid
            cliente = self.modulo.app.test_client()
            # No se inyecta el rol: user_loader lo consulta en PostgreSQL en cada solicitud.
            with cliente.session_transaction() as sesion:
                sesion['_user_id'] = str(uid)
                sesion['_fresh'] = True
            self.clientes[rol] = cliente
        for rol in ('Administrador', 'Cajero'):
            turno = self.filas("SELECT id_turno FROM turno_caja WHERE id_usuario=:u AND estado='ABIERTO' ORDER BY id_turno DESC LIMIT 1",
                               u=self.actores[rol])
            abierto = not turno
            if abierto:
                respuesta = self.formulario(rol, '/caja', accion='abrir', fondo_inicial='0')
                comprobar(respuesta.status_code == 302, 'No fue posible preparar un turno de caja.')
                turno = self.filas("SELECT id_turno FROM turno_caja WHERE id_usuario=:u AND estado='ABIERTO' ORDER BY id_turno DESC LIMIT 1",
                                   u=self.actores[rol])
                comprobar(bool(turno), 'No se creó el turno requerido para vender.')
                self.turnos_nuevos.append((rol, turno[0]['id_turno']))
            self.reporte['preparacion'].append({
                'rol': rol, 'id_usuario': self.actores[rol], 'id_turno': turno[0]['id_turno'],
                'accion': 'Turno de fondo cero abierto y cerrado por esta prueba' if abierto else 'Turno existente reutilizado sin cambiar su saldo ni estado',
            })

    def catalogos_administrador(self):
        rol = 'Administrador'
        r = self.formulario(rol, '/gestion/categorias', nombre='PRUEBA Roles Libreria', descripcion='Datos ficticios de demostracion')
        self.categoria = self.valor("SELECT id_categoria FROM categoria WHERE nombre='PRUEBA Roles Libreria'")
        self.operacion(rol, 'Crear categoría PRUEBA', '/gestion/categorias', r, id_categoria=self.categoria)

        r = self.formulario(rol, '/gestion/proveedores', razon_social='PRUEBA Proveedor Roles',
                            contacto='Contacto de Prueba', correo='proveedor.roles@example.test',
                            direccion='PRUEBA sin valor comercial')
        self.proveedor = self.valor("SELECT id_proveedor FROM proveedor WHERE razon_social='PRUEBA Proveedor Roles'")
        self.operacion(rol, 'Crear proveedor PRUEBA', '/gestion/proveedores', r, id_proveedor=self.proveedor)

        self.formularios_productos = []
        for codigo, nombre, costo, precio in zip(CODIGOS,
                ('PRUEBA Cuaderno Roles', 'PRUEBA Lapicero Roles', 'PRUEBA Carpeta Roles'), ('9', '12', '5'), ('15', '20', '10')):
            datos = dict(id_categoria=str(self.categoria), codigo=codigo, nombre=nombre,
                         marca='Prueba', unidad='Unidad', costo=costo, precio_venta=precio,
                         stock_minimo='5', stock_inicial='100', estado='on')
            r = self.formulario(rol, '/productos/nuevo', **datos)
            pid = self.valor('SELECT id_producto FROM producto WHERE codigo=:codigo', codigo=codigo)
            comprobar(self.stock(pid) == 100, 'El inventario inicial no coincide con la carga de prueba.')
            self.productos.append(pid)
            self.formularios_productos.append(datos)
            self.operacion(rol, f'Crear {nombre} con cien unidades', '/productos/nuevo', r, id_producto=pid, existencia=100)

        ruta = f'/productos/{self.productos[0]}/editar'
        r = self.formulario(rol, ruta, **{**self.formularios_productos[0], 'costo': '9.50', 'precio_venta': '16'})
        producto = self.filas('SELECT costo,precio_venta FROM producto WHERE id_producto=:id', id=self.productos[0])[0]
        comprobar(producto['costo'] == Decimal('9.50') and producto['precio_venta'] == Decimal('16'), 'No se actualizaron el costo y el precio.')
        self.operacion(rol, 'Modificar costo y precio de producto PRUEBA', ruta, r, id_producto=self.productos[0], costo='9.50', precio_venta='16.00')

        self.cliente_admin = self.crear_cliente(rol, 'PRUEBA Cliente Administrador', 'admin.roles@example.test')
        r = self.cobrar(rol, self.cliente_admin, self.productos[1], descuento='1')
        self.venta_admin = self.validar_venta(rol, r, self.cliente_admin, self.productos[1], 1, Decimal('19'))
        self.operacion(rol, 'Registrar venta PRUEBA con descuento autorizado', '/pos/cobrar', r, **self.venta_admin)

    def crear_cliente(self, rol, nombre, correo):
        r = self.formulario(rol, '/gestion/clientes', nombre_razon_social=nombre,
                            correo=correo, direccion_fiscal='PRUEBA sin valor comercial')
        cid = self.valor('SELECT id_cliente FROM cliente WHERE correo=:correo', correo=correo)
        self.operacion(rol, f'Registrar {nombre}', '/gestion/clientes', r, id_cliente=cid)
        return cid

    def validar_venta(self, rol, respuesta, cliente, producto, cantidad, total):
        comprobar(respuesta.status_code == 200 and respuesta.is_json, 'No se pudo registrar una venta de prueba.')
        venta = respuesta.get_json()
        comprobar(venta.get('ok') is True, 'La aplicación rechazó una venta de prueba.')
        fila = self.filas('''SELECT v.id_usuario,v.id_cliente,v.total,v.metodo_pago,d.id_producto,d.cantidad,d.descuento
                             FROM venta v JOIN detalle_venta d USING(id_venta) WHERE v.id_venta=:id''', id=venta['id_venta'])[0]
        comprobar(fila['id_usuario'] == self.actores[rol] and fila['id_cliente'] == cliente
                  and fila['id_producto'] == producto and fila['cantidad'] == cantidad
                  and fila['total'] == total and fila['metodo_pago'] in ('TARJETA', 'TRANSFERENCIA'),
                  'La venta no coincide con el usuario, las cantidades o el importe esperado.')
        self.ventas.append(venta['id_venta'])
        return {'id_venta': venta['id_venta'], 'ticket': venta['numero'], 'total': venta['total'], 'metodo': fila['metodo_pago']}

    def operaciones_cajero(self):
        rol = 'Cajero'
        clientes = [self.crear_cliente(rol, 'PRUEBA Cliente Cajero Alfa', 'cajero.alfa.roles@example.test'),
                    self.crear_cliente(rol, 'PRUEBA Cliente Cajero Beta', 'cajero.beta.roles@example.test')]
        self.clientes_cajero = clientes
        self.ventas_cajero = []
        for i in range(8):
            producto = self.productos[i % 3]
            cantidad = 1 + i % 2
            precio = (Decimal('16'), Decimal('20'), Decimal('10'))[i % 3]
            metodo = ('TRANSFERENCIA', 'TARJETA')[i % 2]
            r = self.cobrar(rol, clientes[i % 2], producto, cantidad, metodo=metodo)
            venta = self.validar_venta(rol, r, clientes[i % 2], producto, cantidad, precio * cantidad)
            self.ventas_cajero.append(venta)
            self.operacion(rol, f'Registrar venta PRUEBA {i + 1} sin descuentos', '/pos/cobrar', r, **venta)

    def operaciones_bodeguero(self):
        rol = 'Bodeguero'
        movimientos = (
            (0, 'ENTRADA', 12), (1, 'ENTRADA', 10), (2, 'ENTRADA', 8), (0, 'ENTRADA', 3),
            (0, 'SALIDA', 2), (1, 'SALIDA', 3), (2, 'SALIDA', 1),
        )
        for indice, tipo, cantidad in movimientos:
            pid = self.productos[indice]
            antes = self.stock(pid)
            motivo = f'PRUEBA Roles {tipo.lower()} de demostracion'
            r = self.formulario(rol, '/inventario/movimiento', id_producto=pid, tipo=tipo, cantidad=cantidad, motivo=motivo)
            despues = self.stock(pid)
            movimiento = self.filas('''SELECT id_movimiento_inventario,id_usuario,cantidad,tipo FROM movimiento_inventario
                                      WHERE id_producto=:p ORDER BY id_movimiento_inventario DESC LIMIT 1''', p=pid)[0]
            esperado = antes + (cantidad if tipo == 'ENTRADA' else -cantidad)
            comprobar(despues == esperado and movimiento['id_usuario'] == self.actores[rol]
                      and movimiento['cantidad'] == cantidad and movimiento['tipo'] == tipo,
                      'No se guardó correctamente el movimiento del bodeguero.')
            self.operacion(rol, f'Registrar {tipo.lower()} PRUEBA de {cantidad} unidades', '/inventario/movimiento', r,
                           id_movimiento=movimiento['id_movimiento_inventario'], id_producto=pid,
                           cantidad=cantidad, existencia_antes=antes, existencia_despues=despues)
        for pid, diferencia, destino in zip(self.productos, (2, -1, -3), ('aprobar', 'rechazar', 'pendiente')):
            antes = self.stock(pid)
            motivo = f'PRUEBA Roles conteo para {destino}'
            r = self.formulario(rol, '/inventario/conteos', id_producto=pid, cantidad_contada=antes + diferencia, motivo=motivo)
            solicitudes = self.filas('SELECT * FROM solicitud_ajuste WHERE id_producto=:p ORDER BY id_solicitud DESC LIMIT 1', p=pid)
            comprobar(bool(solicitudes), 'No se creó la solicitud de ajuste; se revierte toda la demostración.')
            solicitud = solicitudes[0]
            comprobar(solicitud['id_solicitante'] == self.actores[rol] and solicitud['estado'] == 'PENDIENTE'
                      and solicitud['diferencia'] == diferencia, 'No se creó la solicitud de ajuste pendiente.')
            self.verificar(f'Conteo para {destino} no cambia existencias antes de aprobar', self.stock(pid) == antes,
                           id_solicitud=solicitud['id_solicitud'], existencia=antes)
            self.solicitudes.append(solicitud['id_solicitud'])
            self.operacion(rol, f'Registrar conteo PRUEBA para {destino}', '/inventario/conteos', r,
                           id_solicitud=solicitud['id_solicitud'], id_producto=pid,
                           existencia_sistema=antes, cantidad_contada=antes + diferencia, diferencia=diferencia)

    def resolver_ajustes(self):
        rol = 'Administrador'
        pid = self.productos[0]
        antes = self.stock(pid)
        ruta = f'/inventario/ajustes/{self.solicitudes[0]}/aprobar'
        r = self.formulario(rol, ruta)
        solicitud = self.filas('SELECT * FROM solicitud_ajuste WHERE id_solicitud=:id', id=self.solicitudes[0])[0]
        comprobar(solicitud['estado'] == 'APROBADA' and solicitud['id_aprobador'] == self.actores[rol]
                  and solicitud['id_movimiento'] is not None and self.stock(pid) == antes + 2,
                  'No se aplicó correctamente el ajuste aprobado.')
        self.operacion(rol, 'Aprobar y aplicar una solicitud de ajuste PRUEBA', ruta, r,
                       id_solicitud=self.solicitudes[0], id_movimiento=solicitud['id_movimiento'], existencia_despues=self.stock(pid))
        antes = self.stock(self.productos[1])
        ruta = f'/inventario/ajustes/{self.solicitudes[1]}/rechazar'
        r = self.formulario(rol, ruta, motivo='PRUEBA rechazo de demostracion')
        solicitud = self.filas('SELECT * FROM solicitud_ajuste WHERE id_solicitud=:id', id=self.solicitudes[1])[0]
        comprobar(solicitud['estado'] == 'RECHAZADA' and solicitud['id_aprobador'] == self.actores[rol]
                  and solicitud['id_movimiento'] is None and self.stock(self.productos[1]) == antes,
                  'El rechazo debe conservar las existencias.')
        self.operacion(rol, 'Rechazar solicitud de ajuste PRUEBA sin cambiar existencias', ruta, r,
                       id_solicitud=self.solicitudes[1], existencia_despues=antes)

    def comprobar_permisos(self):
        cajero, bodeguero, admin = 'Cajero', 'Bodeguero', 'Administrador'
        venta_ajena = self.venta_admin['id_venta']
        self.denegado('Cajero no consulta el comprobante de la venta del administrador',
                      self.clientes[cajero].get(f'/ventas/{venta_ajena}/ticket'))
        self.denegado('Cajero no modifica precios por URL directa',
                      self.formulario(cajero, f'/productos/{self.productos[0]}/editar', precio_venta='0.01'))
        self.denegado('Cajero no inyecta un descuento desde JSON',
                      self.cobrar(cajero, self.clientes_cajero[0], self.productos[0], descuento='1'))
        self.denegado('Cajero no registra movimientos de inventario',
                      self.formulario(cajero, '/inventario/movimiento', id_producto=self.productos[0], tipo='ENTRADA', cantidad=1, motivo='PRUEBA prohibida'))
        self.denegado('Cajero no elimina productos por URL directa',
                      self.formulario(cajero, f'/productos/{self.productos[0]}/eliminar'))
        self.denegado('Cajero no anula ventas',
                      self.formulario(cajero, f'/ventas/{self.ventas_cajero[0]["id_venta"]}/anular', motivo='PRUEBA prohibida'))
        for rol in (cajero, bodeguero):
            self.denegado(f'{rol} no administra usuarios', self.clientes[rol].get('/usuarios'))
        self.denegado('Bodeguero no cobra ventas', self.cobrar(bodeguero, self.cliente_admin, self.productos[0]))
        self.denegado('Bodeguero no consulta clientes', self.clientes[bodeguero].get('/gestion/clientes'))
        self.denegado('Bodeguero no aprueba ajustes',
                      self.formulario(bodeguero, f'/inventario/ajustes/{self.solicitudes[2]}/aprobar'))

        for venta in self.ventas_cajero:
            respuesta = self.clientes[cajero].get(f'/ventas/{venta["id_venta"]}/ticket')
            self.verificar(f'Cajero consulta su comprobante {venta["ticket"]}', respuesta.status_code == 200)
        html = self.clientes[cajero].get('/ventas').get_data(as_text=True)
        self.verificar('Listado del cajero incluye sus ocho ventas PRUEBA y excluye la venta ajena',
                       all(v['ticket'] in html for v in self.ventas_cajero) and self.venta_admin['ticket'] not in html)
        for ruta in ('/reportes/excel', '/reportes/ventas/excel'):
            respuesta = self.clientes[cajero].get(ruta)
            comprobar(respuesta.status_code == 200, 'No se pudo exportar el reporte del cajero.')
            libro = load_workbook(BytesIO(respuesta.data), read_only=True, data_only=True)
            try:
                filas = list(libro.active.iter_rows(values_only=True))
            finally:
                libro.close()
            tickets = {f[0] for f in filas[1:]}
            self.verificar(f'Exportación {ruta} respeta ventas propias y omite costos y ganancias',
                           all(v['ticket'] in tickets for v in self.ventas_cajero)
                           and self.venta_admin['ticket'] not in tickets
                           and not any('costo' in str(c).lower() or 'ganancia' in str(c).lower() for c in filas[0]), http=200)

        respuesta = self.clientes[bodeguero].get('/inventario/excel')
        comprobar(respuesta.status_code == 200, 'No se pudo exportar inventario del bodeguero.')
        libro = load_workbook(BytesIO(respuesta.data), read_only=True)
        try:
            encabezados = next(libro.active.iter_rows(values_only=True))
        finally:
            libro.close()
        self.verificar('Exportación de inventario del bodeguero no incluye costos',
                       not any('costo' in str(c).lower() for c in encabezados))
        self.verificar('Administrador consulta reportes generales', self.clientes[admin].get('/reportes').status_code == 200)

        pid = self.productos[0]
        antes, movimientos = self.stock(pid), self.valor('SELECT COUNT(*) FROM movimiento_inventario')
        self.formulario(admin, f'/inventario/ajustes/{self.solicitudes[0]}/aprobar')
        self.verificar('Un ajuste aprobado no se aplica dos veces', self.stock(pid) == antes
                       and self.valor('SELECT COUNT(*) FROM movimiento_inventario') == movimientos)
        self.formulario(bodeguero, '/inventario/movimiento', id_producto=pid, tipo='SALIDA', cantidad=antes + 1, motivo='PRUEBA stock insuficiente')
        self.verificar('Salida superior al stock no crea movimientos ni existencias negativas', self.stock(pid) == antes
                       and self.valor('SELECT COUNT(*) FROM movimiento_inventario') == movimientos)
        self.formulario(admin, f'/productos/{pid}/eliminar')
        self.verificar('Un producto con ventas e historial no se elimina',
                       self.valor('SELECT COUNT(*) FROM producto WHERE id_producto=:id', id=pid) == 1
                       and self.valor('SELECT COUNT(*) FROM movimiento_inventario') == movimientos)
        pendiente = self.filas('SELECT * FROM solicitud_ajuste WHERE id_solicitud=:id', id=self.solicitudes[2])[0]
        self.verificar('La tercera solicitud permanece pendiente sin modificar stock',
                       pendiente['estado'] == 'PENDIENTE' and pendiente['id_movimiento'] is None
                       and self.stock(self.productos[2]) == pendiente['existencia_sistema'])
        uid = self.actores[admin]
        self.formulario(admin, f'/usuarios/{uid}/estado', estado='desactivar')
        self.verificar('Administrador no desactiva su propia cuenta',
                       self.valor('SELECT estado FROM usuario WHERE id_usuario=:id', id=uid) is True)
        cantidad_admins = self.valor("SELECT COUNT(*) FROM usuario u JOIN rol r USING(id_rol) WHERE u.estado=TRUE AND r.nombre='Administrador'")
        if cantidad_admins == 1:
            rol_cajero = self.valor('SELECT id_rol FROM usuario WHERE id_usuario=:id', id=self.actores[cajero])
            self.formulario(admin, f'/usuarios/{uid}/rol', id_rol=rol_cajero)
            self.verificar('No se quita el rol al último administrador activo',
                           self.valor('SELECT r.nombre FROM usuario u JOIN rol r USING(id_rol) WHERE u.id_usuario=:id', id=uid) == admin)

    def terminar(self):
        for rol, tid in self.turnos_nuevos:
            self.formulario(rol, '/caja', accion='cerrar', efectivo_contado='0')
            comprobar(self.valor('SELECT estado FROM turno_caja WHERE id_turno=:id', id=tid) == 'CERRADO',
                      'No se pudo cerrar el turno creado por la demostración.')
        self.verificar('No se crearon usuarios ni roles', self.valor('SELECT COUNT(*) FROM usuario') == self.usuarios_antes
                       and self.valor('SELECT COUNT(*) FROM rol') == self.roles_antes)
        self.verificar('No se modificó el efectivo ni se crearon movimientos de caja',
                       self.valor('SELECT COUNT(*) FROM movimiento_caja') == self.caja_antes)
        totales = dict(Counter(o['rol'] for o in self.reporte['operaciones']))
        comprobar(totales == {r: 10 for r in ROLES}, 'El lote debe contener exactamente diez operaciones por cada rol.')
        self.reporte['total_operaciones'] = len(self.reporte['operaciones'])
        self.reporte['operaciones_por_rol'] = totales
        self.reporte['usuarios_responsables'] = self.actores
        self.reporte['registros'] = {
            'categoria': self.categoria, 'proveedor': self.proveedor,
            'productos': [{'id_producto': pid, 'codigo': codigo, 'existencia_final': self.stock(pid)}
                          for pid, codigo in zip(self.productos, CODIGOS)],
            'clientes': [self.cliente_admin, *self.clientes_cajero], 'ventas': self.ventas,
            'solicitudes': self.filas('''SELECT id_solicitud,id_producto,estado,existencia_sistema,cantidad_contada,
                                        diferencia,id_solicitante,id_aprobador,id_movimiento FROM solicitud_ajuste
                                        WHERE id_solicitud IN (:a,:b,:c) ORDER BY id_solicitud''',
                                      a=self.solicitudes[0], b=self.solicitudes[1], c=self.solicitudes[2]),
        }
        return self.reporte


def ejecutar_demostracion(modulo, engine, *, aplicar=False):
    """Una transacción exterior conserva todo o revierte todo, incluidos commits HTTP.

    Las sesiones de cada solicitud usan savepoints sobre la misma conexión. Este
    comando debe ejecutarse en un proceso independiente del servidor web.
    """
    comprobar(engine.dialect.name == 'postgresql', 'Esta demostración requiere PostgreSQL.')
    comprobar(modulo.app.config.get('WTF_CSRF_ENABLED', True), 'La demostración requiere CSRF habilitado.')
    with modulo.app.app_context():
        modulo.db.session.remove()
    sesion_original = modulo.db.session
    with engine.connect() as conexion:
        transaccion = conexion.begin()
        sesiones = scoped_session(sessionmaker(bind=conexion, join_transaction_mode='create_savepoint'))
        modulo.db.session = sesiones
        try:
            # El bloqueo pertenece a la transacción exterior: el rollback de una
            # solicitud HTTP no debe liberarlo antes de terminar la demostración.
            conexion.exec_driver_sql("SET LOCAL lock_timeout='5s'")
            conexion.exec_driver_sql("SET LOCAL statement_timeout='15s'")
            comprobar(conexion.execute(text("SELECT pg_try_advisory_xact_lock(hashtextextended('puntocom:demo_roles',0))")).scalar_one(),
                      'Otra carga de demostración está en curso.')
            demo = Demostracion(modulo)
            demo.preparar()
            demo.catalogos_administrador()
            demo.operaciones_cajero()
            demo.operaciones_bodeguero()
            demo.resolver_ajustes()
            demo.comprobar_permisos()
            reporte = demo.terminar()
            sesiones.remove()
            if aplicar:
                transaccion.commit()
            else:
                transaccion.rollback()
            reporte['aplicado'] = aplicar
            reporte['base'] = engine.url.database
            return reporte
        except BaseException:
            sesiones.remove()
            if transaccion.is_active:
                transaccion.rollback()
            raise
        finally:
            modulo.db.session = sesion_original


def guardar_informe(reporte):
    carpeta = ROOT / 'documentacion'
    carpeta.mkdir(exist_ok=True)
    sufijo = '' if reporte['aplicado'] else '_simulacion'
    nombre = f'Prueba_30_operaciones_roles{sufijo}.md'
    estado = 'GUARDADAS EN LA BASE ACTUAL' if reporte['aplicado'] else 'SIMULADAS Y REVERTIDAS'
    lineas = [
        '# PUNTO COM: treinta operaciones de prueba por roles', '',
        f"Fecha: {reporte['fecha']}. Base: `{reporte['base']}`. Estado: **{estado}**.", '',
        'Se reutilizaron tres usuarios existentes. No se crearon cuentas ni roles, ni se cambiaron contraseñas.', '',
        'Las operaciones se ejecutaron por las rutas HTTP reales, con sesiones autenticadas asignadas desde el servidor '
        'y tokens CSRF válidos. Esta carga no comprueba contraseñas; la suite de autorización prueba el inicio de sesión.', '',
        'Cada formulario se comprobó consultando sus registros y existencias en PostgreSQL; una redirección HTTP 302 '
        'por sí sola no se consideró un éxito. Todo el lote se confirma en una sola transacción después de validar los resultados.', '',
        'Son datos ficticios identificados por PRUEBA en productos, categoría, proveedor y clientes. Las nueve ventas '
        'se reconocen por sus clientes/productos y por los tickets indicados abajo. **Sí aparecen en los reportes de ventas '
        'y existencias habituales.** Se usaron TARJETA/TRANSFERENCIA sin conexiones bancarias ni movimientos de efectivo.', '',
        'La apertura/cierre de turnos de fondo cero, cuando fueron necesarios, se registra como preparación adicional. '
        'Los turnos previamente abiertos conservan su estado y saldo. No se borró historial existente.', '',
    ]
    for rol in ROLES:
        lineas += [f'## {rol}: 10 operaciones', '', '| Nº | Operación | Ruta | HTTP | Evidencia |', '| --- | --- | --- | --- | --- |']
        for op in reporte['operaciones']:
            if op['rol'] == rol:
                evidencia = ', '.join(f'{k}: {v}' for k, v in op.items()
                                      if k not in ('rol', 'numero', 'nombre', 'ruta', 'http', 'resultado'))
                lineas.append(f"| {op['numero']} | {op['nombre']} | `{op['ruta']}` | {op['http']} | {evidencia} |")
        lineas.append('')
    lineas += ['## Comprobaciones adicionales', '']
    for prueba in reporte['comprobaciones']:
        lineas.append(f"- CORRECTO: {prueba['nombre']}.")
    lineas += ['', '## Preparación de caja', '']
    for preparacion in reporte['preparacion']:
        lineas.append(f"- {preparacion['rol']}, usuario {preparacion['id_usuario']}, turno {preparacion['id_turno']}: {preparacion['accion']}.")
    lineas += ['', '## Cómo revisar los datos', '',
               '1. Inicia sesión con tu administrador habitual. En Productos busca `PRUEBA-ROLES`; en Ventas busca un ticket del listado. '
               'En Inventario → Ajustes encontrarás una solicitud aprobada, una rechazada y una pendiente.',
               '2. Entra con el cajero existente: podrá consultar sus ocho ventas de prueba y sus comprobantes. '
               'El comprobante de la venta del administrador debe devolver HTTP 403.',
               '3. Entra con el bodeguero existente: podrá revisar los movimientos y el conteo pendiente; '
               'no podrá aprobarlo, cobrar ni administrar usuarios.', '',
               '## Repetición y alcance', '',
               'La carga rechaza una segunda ejecución si existen los códigos PRUEBA-ROLES-A/B/C, para evitar duplicados. '
               'No incluye un borrado automático: las ventas y movimientos ya conservados son historial.', '',
               'Para validar sin conservar otro lote: `python demostracion_roles.py` (antes de aplicar esta carga). '
               'Para conservarlo: `python demostracion_roles.py --aplicar`. Una simulación revierte los registros, pero '
               'PostgreSQL puede reservar números de secuencia que no se reutilizan.', '',
               'Estas treinta operaciones cubren una demostración representativa. Las demás reglas se verifican con '
               '`python tests/postgres/ejecutar.py`; editar la configuración general sigue pendiente de implementación.', '']
    (carpeta / nombre).write_text('\n'.join(lineas), encoding='utf-8')
    (carpeta / f'resultado_30_operaciones_roles{sufijo}.json').write_text(
        json.dumps(reporte, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return carpeta / nombre


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--aplicar', action='store_true', help='Conservar los datos ficticios en la base configurada en .env.')
    argumentos = parser.parse_args()
    import app as modulo
    try:
        with modulo.app.app_context():
            engine = modulo.db.engine
        reporte = ejecutar_demostracion(modulo, engine, aplicar=argumentos.aplicar)
    except ErrorDemostracion as error:
        print(f'No se guardó el lote: {error}')
        return 1
    except Exception as error:
        # No imprimir excepciones SQL, parámetros ni direcciones con credenciales.
        print(f'No se guardó el lote. Error de infraestructura: {type(error).__name__}.')
        return 1
    print(f"30 operaciones correctas: 10 por rol; {len(reporte['comprobaciones'])} comprobaciones adicionales correctas.")
    print('Datos guardados.' if reporte['aplicado'] else 'Simulación terminada; registros revertidos.')
    try:
        informe = guardar_informe(reporte)
    except Exception as error:
        print(f'No se pudo guardar el informe: {type(error).__name__}.')
        if reporte['aplicado']:
            print('El lote SÍ quedó guardado. No repitas la carga; busca PRUEBA-ROLES en Productos.')
            print('Ventas conservadas:', ', '.join(str(vid) for vid in reporte['registros']['ventas']))
        return 2
    print(f'Informe: {informe}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
