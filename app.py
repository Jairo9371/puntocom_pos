
from datetime import datetime, date, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from flask import send_file
from sqlalchemy import text
import os
import re
import secrets
from datetime import datetime, date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from functools import wraps
from io import BytesIO
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, url_for, abort
from flask_login import LoginManager, UserMixin, current_user, login_required, login_user, logout_user
from flask_sqlalchemy import SQLAlchemy
from flask_wtf.csrf import CSRFProtect
from sqlalchemy import text
from sqlalchemy.engine import URL
from sqlalchemy.exc import IntegrityError
from werkzeug.security import check_password_hash, generate_password_hash
from flask_login import current_user

load_dotenv()
app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY') or ('dev-only-' + secrets.token_hex(32))
app.config['SQLALCHEMY_DATABASE_URI'] = URL.create(
    'postgresql+psycopg2', username=os.getenv('DB_USER', 'postgres'),
    password=os.getenv('DB_PASSWORD', ''), host=os.getenv('DB_HOST', '127.0.0.1'),
    port=int(os.getenv('DB_PORT', '5433')), database=os.getenv('DB_NAME', 'puntocom_pos'))
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['MAX_CONTENT_LENGTH'] = 4 * 1024 * 1024
csrf = CSRFProtect(app)
db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
TZ = ZoneInfo('America/Guatemala')

class Usuario(UserMixin):
    def __init__(self, row):
        self.id = str(row['id_usuario'])
        self.nombre = row['nombre']
        self.correo = row['correo']
        self.rol = row['rol']

@login_manager.user_loader
def load_user(user_id):
    try:
        row = one('''SELECT u.id_usuario,u.nombre,u.correo,r.nombre AS rol
                     FROM usuario u JOIN rol r USING(id_rol)
                     WHERE u.id_usuario=:id AND u.estado=TRUE''', id=int(user_id))
        return Usuario(row) if row else None
    except (ValueError, Exception):
        return None

def rows(sql, **params):
    return [dict(r) for r in db.session.execute(text(sql), params).mappings().all()]

def one(sql, **params):
    row = db.session.execute(text(sql), params).mappings().first()
    return dict(row) if row else None

def value(sql, **params):
    return db.session.execute(text(sql), params).scalar()

def money(x):
    return Decimal(str(x or 0)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def numeric(key, minimum=0):
    try:
        result = money(request.form.get(key, '0'))
        if result < minimum: raise ValueError()
        return result
    except (ValueError, InvalidOperation):
        raise ValueError(f'Valor inválido para {key}.')

def number(key, minimum=0):
    try:
        n = int(request.form.get(key, ''))
        if n < minimum: raise ValueError()
        return n
    except (ValueError, TypeError):
        raise ValueError(f'Cantidad inválida para {key}.')

def required(key):
    x = request.form.get(key, '').strip()
    if not x: raise ValueError(f'El campo {key} es obligatorio.')
    return x

def role_required(*allowed):
    def decorate(fn):
        @wraps(fn)
        @login_required
        def wrapper(*args, **kwargs):
            if current_user.rol not in allowed:
                abort(403)
            return fn(*args, **kwargs)
        return wrapper
    return decorate

ADMIN = ('Administrador',)
INV = ('Administrador', 'Bodeguero', 'Encargado de inventario')
VENTA = ('Administrador', 'Cajero')

@app.template_filter('q')
def quetzales(v):
    return 'Q {:,.2f}'.format(money(v))

@app.template_filter('fecha')
def fecha(v):
    return v.strftime('%d/%m/%Y %H:%M') if v else '-'

@app.context_processor
def global_vars():
    alertas = []
    total_alertas = 0

    if current_user.is_authenticated:
        alertas = rows("""
            SELECT
                vi.codigo,
                vi.nombre,
                vi.existencia_actual,
                vi.stock_minimo
            FROM vista_inventario vi
            JOIN producto p
                ON p.id_producto = vi.id_producto
            WHERE p.estado = TRUE
              AND vi.existencia_actual <= vi.stock_minimo
            ORDER BY vi.existencia_actual ASC
            LIMIT 20
        """)

        total_alertas = len(alertas)

    return {
        'shop_name': 'PUNTO COM',
        'clock_today': datetime.now(TZ).strftime('%d/%m/%Y'),
        'total_alertas': total_alertas,
        'alertas': alertas
    }


@app.errorhandler(403)
def no_permisos(_):
    return render_template('message.html', titulo='Acceso denegado', mensaje='No tienes permisos para realizar esta operación.'), 403

@app.errorhandler(404)
def no_existe(_):
    return render_template('message.html', titulo='No encontrado', mensaje='El registro o la página no existe.'), 404

@app.route('/login', methods=['GET','POST'])
def login():
    if current_user.is_authenticated: return redirect(url_for('dashboard'))
    if request.method == 'POST':
        correo = request.form.get('correo','').strip().lower()
        row = one('''SELECT u.id_usuario,u.nombre,u.correo,u.hash_contrasena,r.nombre AS rol
                     FROM usuario u JOIN rol r USING(id_rol)
                     WHERE LOWER(u.correo)=:correo AND u.estado=TRUE''', correo=correo)
        if row and check_password_hash(row['hash_contrasena'], request.form.get('password','')):
            login_user(Usuario(row), remember=bool(request.form.get('remember')))
            return redirect(url_for('dashboard'))
        flash('Correo o contraseña incorrectos.', 'danger')
    return render_template('login.html')

@app.post('/logout')
@login_required
def logout():
    logout_user()
    flash('Sesión cerrada.', 'success')
    return redirect(url_for('login'))


@app.get('/')
@login_required
def dashboard():

    # Estadísticas principales
    stats = one("""
        SELECT
            (SELECT COALESCE(SUM(total), 0)
             FROM venta
             WHERE fecha_hora::date = CURRENT_DATE
             AND estado = 'COMPLETADA') AS ventas,

            (SELECT COUNT(*)
             FROM venta
             WHERE fecha_hora::date = CURRENT_DATE
             AND estado = 'COMPLETADA') AS nventas,

            (SELECT COUNT(*)
             FROM producto
             WHERE estado = TRUE) AS productos,

            (SELECT COUNT(*)
             FROM vista_inventario
             WHERE existencia_actual <= stock_minimo
             AND existencia_actual > 0) AS bajos,

            (SELECT COUNT(*)
             FROM cliente) AS clientes,

            (SELECT COALESCE(SUM(
                (dv.precio_unitario * dv.cantidad - dv.descuento)
                - (dv.costo_unitario_historico * dv.cantidad)
             ), 0)
             FROM detalle_venta dv
             JOIN venta v USING(id_venta)
             WHERE v.fecha_hora::date = CURRENT_DATE
             AND v.estado = 'COMPLETADA') AS ganancia
    """)

    # Últimas ventas realizadas
    recientes = rows("""
        SELECT
            v.numero_ticket,
            v.fecha_hora,
            v.total,
            c.nombre_razon_social
        FROM venta v
        JOIN cliente c USING(id_cliente)
        ORDER BY v.id_venta DESC
        LIMIT 8
    """)

    # Productos más vendidos
    top = rows("""
        SELECT
            p.nombre,
            SUM(d.cantidad) AS cantidad
        FROM detalle_venta d
        JOIN venta v USING(id_venta)
        JOIN producto p USING(id_producto)
        WHERE v.estado = 'COMPLETADA'
        GROUP BY p.id_producto, p.nombre
        ORDER BY cantidad DESC
        LIMIT 5
    """)

    # Ventas de los últimos 7 días
    dias = rows("""
        SELECT
            d::date AS fecha,
            COALESCE(SUM(v.total), 0) AS total
        FROM generate_series(
            CURRENT_DATE - INTERVAL '6 days',
            CURRENT_DATE,
            INTERVAL '1 day'
        ) d
        LEFT JOIN venta v
            ON v.fecha_hora::date = d::date
            AND v.estado = 'COMPLETADA'
        GROUP BY d
        ORDER BY d
    """)
    
    # Ventas mensuales del año actual
    ventas_mensuales = rows("""
        SELECT
            m.mes,
            COALESCE(SUM(v.total), 0) AS total
        FROM generate_series(1, 12) AS m(mes)
        LEFT JOIN venta v
            ON EXTRACT(MONTH FROM v.fecha_hora) = m.mes
            AND EXTRACT(YEAR FROM v.fecha_hora)
                = EXTRACT(YEAR FROM CURRENT_DATE)
            AND v.estado = 'COMPLETADA'
        GROUP BY m.mes
        ORDER BY m.mes
    """)

    meses = [
        "Enero", "Febrero", "Marzo", "Abril",
        "Mayo", "Junio", "Julio", "Agosto",
        "Septiembre", "Octubre", "Noviembre",
        "Diciembre"
    ]

    totales_mensuales = [
        float(fila["total"])
        for fila in ventas_mensuales
    ]

    # Enviar los resultados a la plantilla HTML
    return render_template(
        'dashboard.html',
        stats=stats,
        recientes=recientes,
        top=top,
        dias=dias,
        meses=meses,
        totales_mensuales=totales_mensuales
    )

@app.get('/productos')
@login_required
def productos():
    buscar=request.args.get('q','').strip()
    categoria=request.args.get('categoria','')
    filtros={'q':f'%{buscar}%'}
    cond="(p.codigo ILIKE :q OR p.codigo_barras ILIKE :q OR p.nombre ILIKE :q)"
    if categoria.isdigit(): cond+=' AND p.id_categoria=:categoria'; filtros['categoria']=int(categoria)
    data=rows(f'''SELECT p.*,c.nombre AS categoria,vi.existencia_actual
                   FROM producto p JOIN categoria c USING(id_categoria)
                   JOIN vista_inventario vi USING(id_producto)
                   WHERE {cond} ORDER BY p.codigo''',**filtros)
    return render_template('productos.html', data=data, categorias=rows('SELECT * FROM categoria ORDER BY nombre'))

@app.route('/productos/nuevo',methods=['GET','POST'])
@app.route('/productos/<int:pid>/editar',methods=['GET','POST'])
@role_required(*INV)
def producto_form(pid=None):
    producto=one('SELECT * FROM producto WHERE id_producto=:id',id=pid) if pid else None
    if pid and not producto: abort(404)
    if request.method=='POST':
        try:
            payload={'categoria':number('id_categoria',1),'codigo':required('codigo'),
                     'barras':request.form.get('codigo_barras','').strip() or None,
                     'nombre':required('nombre'),
                     'marca':solo_letras(request.form.get('marca',''), 'Marca', obligatorio=False),
                     'unidad':required('unidad'),'costo':numeric('costo'),
                     'precio':numeric('precio_venta'),'minimo':number('stock_minimo'),
                     'activo': bool(request.form.get('estado'))}
            if pid:
                db.session.execute(text('''UPDATE producto SET id_categoria=:categoria,codigo=:codigo,
                  codigo_barras=:barras,nombre=:nombre,marca=:marca,unidad=:unidad,costo=:costo,
                  precio_venta=:precio,stock_minimo=:minimo,estado=:activo WHERE id_producto=:id'''),
                  {**payload,'id':pid})
            else:
                pid_nuevo=value('''INSERT INTO producto(id_categoria,codigo,codigo_barras,nombre,marca,unidad,
                          costo,precio_venta,stock_minimo,estado)
                          VALUES(:categoria,:codigo,:barras,:nombre,:marca,:unidad,:costo,:precio,:minimo,:activo)
                          RETURNING id_producto''',**payload)
                stock=number('stock_inicial')
                if stock:
                    db.session.execute(text('''INSERT INTO movimiento_inventario
                     (id_producto,tipo,cantidad,costo_unitario,motivo)
                     VALUES(:id,'ENTRADA',:cantidad,:costo,'Inventario inicial')'''),
                     {'id':pid_nuevo,'cantidad':stock,'costo':payload['costo']})
            db.session.commit(); flash('Producto guardado correctamente.','success')
            return redirect(url_for('productos'))
        except (ValueError,IntegrityError) as e:
            db.session.rollback(); flash(str(e) if isinstance(e, ValueError) else 'Revisa la categoría o los códigos duplicados.', 'danger')
        except Exception:
            db.session.rollback(); app.logger.exception('Guardar producto'); flash('Error al guardar producto.','danger')
    return render_template('producto_form.html',item=producto,categorias=rows('SELECT * FROM categoria ORDER BY nombre'))

@app.post('/productos/<int:pid>/estado')
@role_required(*INV)
def producto_estado(pid):
    db.session.execute(text('UPDATE producto SET estado=NOT estado WHERE id_producto=:id'),{'id':pid})
    db.session.commit(); flash('Estado del producto actualizado.','success'); return redirect(url_for('productos'))

@app.get('/inventario')
@login_required
def inventario():
    return render_template('inventario.html',data=rows('''SELECT vi.*,p.costo,c.nombre AS categoria
      FROM vista_inventario vi JOIN producto p USING(id_producto) JOIN categoria c USING(id_categoria)
      ORDER BY vi.codigo'''),movimientos=rows('''SELECT mi.*,p.codigo,p.nombre FROM movimiento_inventario mi
      JOIN producto p USING(id_producto) ORDER BY mi.id_movimiento_inventario DESC LIMIT 80'''))

@app.post('/inventario/movimiento')
@role_required(*INV)
def inventario_movimiento():
    try:
        pid=number('id_producto',1); cant=number('cantidad',1); tipo=required('tipo')
        if tipo not in ('ENTRADA','SALIDA'): raise ValueError('Tipo de movimiento inválido.')
        motivo=solo_letras(required('motivo'), 'Motivo de inventario')
        # Adquirir bloqueo antes de leer existencias y actualizar.
        p=one('SELECT id_producto,costo FROM producto WHERE id_producto=:id FOR UPDATE',id=pid)
        if not p: raise ValueError('Producto no existe.')
        saldo=value('SELECT existencia_actual FROM vista_inventario WHERE id_producto=:id',id=pid)
        if tipo=='SALIDA' and cant>saldo: raise ValueError(f'Existencia insuficiente: {saldo}.')
        db.session.execute(text('''INSERT INTO movimiento_inventario(id_producto,tipo,cantidad,costo_unitario,motivo)
          VALUES(:id,:tipo,:cantidad,:costo,:motivo)'''),{'id':pid,'tipo':tipo,'cantidad':cant,'costo':p['costo'],'motivo':motivo})
        db.session.commit(); flash('Movimiento guardado.','success')
    except (ValueError, IntegrityError) as e:
        db.session.rollback(); flash(str(e) if isinstance(e,ValueError) else 'Datos no válidos.','danger')
    except Exception:
        db.session.rollback(); app.logger.exception('Inventario'); flash('No fue posible actualizar inventario.','danger')
    return redirect(url_for('inventario'))

# Reglas solicitadas: letras, espacios, tildes y eñe.
# No se aceptan números, puntos ni guiones en los campos indicados.
RE_SOLO_LETRAS = re.compile(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+(?: [A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+)*")

def solo_letras(valor, etiqueta, obligatorio=True):
    valor = (valor or '').strip()
    if not valor and not obligatorio:
        return valor
    if not RE_SOLO_LETRAS.fullmatch(valor):
        raise ValueError(f'{etiqueta}: solo se permiten letras y espacios.')
    return valor


# Validaciones del backend para los formularios administrativos.
# HTML por sí solo no es una protección suficiente.
RE_NOMBRE_PERSONA = re.compile(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+(?:[ .'-][A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+)*")
RE_TELEFONO_GT = re.compile(r"[0-9]{8}")
RE_DPI = re.compile(r"[0-9]{13}")
RE_NIT = re.compile(r"[0-9]{1,12}-?[0-9K]")
RE_CORREO = re.compile(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?)+")


def validar_datos_gestion(modulo, data):
    """Normaliza campos y valida según la entidad antes de INSERT/UPDATE."""
    # Evitar valores que excedan los límites del esquema actual.
    limites = {
        'categorias': {'nombre': 100},
        'clientes': {'nombre_razon_social': 150, 'nit_dpi': 30, 'telefono': 20, 'correo': 150},
        'proveedores': {'razon_social': 150, 'nit': 30, 'contacto': 100, 'telefono': 20, 'correo': 150},
    }
    for campo, limite in limites.get(modulo, {}).items():
        valor = data.get(campo)
        if valor is not None and len(valor) > limite:
            raise ValueError(f'El campo {campo.replace("_", " ")} supera {limite} caracteres.')

    if modulo == 'categorias':
        data['nombre'] = solo_letras(data['nombre'], 'Nombre de categoría')
        data['descripcion'] = solo_letras(data.get('descripcion'), 'Descripción de categoría', obligatorio=False)

    elif modulo in ('clientes', 'proveedores'):
        campo_nombre = 'nombre_razon_social' if modulo == 'clientes' else 'razon_social'
        nombre = data[campo_nombre]
        if not nombre:
            raise ValueError('El nombre o razón social es obligatorio.')
        if modulo == 'clientes':
            data[campo_nombre] = solo_letras(nombre, 'Nombre de cliente')
        else:
            data[campo_nombre] = nombre  # Se admiten números y signos en razones sociales.

        if modulo == 'proveedores':
            data['contacto'] = solo_letras(data.get('contacto'), 'Nombre del contacto', obligatorio=False)

        telefono = data.get('telefono')
        if telefono and not RE_TELEFONO_GT.fullmatch(telefono):
            raise ValueError('El teléfono debe contener exactamente 8 números.')

        correo = data.get('correo')
        if correo:
            correo = correo.lower()
            if not RE_CORREO.fullmatch(correo):
                raise ValueError('Introduce un correo electrónico válido.')
            data['correo'] = correo

        campo_nit = 'nit_dpi' if modulo == 'clientes' else 'nit'
        nit = data.get(campo_nit)
        if nit:
            nit = nit.upper().replace(' ', '')
            # Consumidor Final es un registro especial, no un NIT normal.
            if modulo == 'clientes' and nit == 'CF':
                if nombre.casefold() != 'consumidor final':
                    raise ValueError('El identificador CF está reservado para Consumidor Final.')
            elif not (RE_DPI.fullmatch(nit) if modulo == 'clientes' else False) and not RE_NIT.fullmatch(nit):
                raise ValueError('Ingrese un NIT válido en formato o DPI de 13 dígitos.')
            data[campo_nit] = nit

    return data


# CRUD sencillo de categorías, clientes y proveedores usando formularios genéricos.
CRUD = {
 'categorias': {'tabla':'categoria','pk':'id_categoria','campos':['nombre','descripcion'],'titulo':'Categorías','roles':INV},
 'clientes': {'tabla':'cliente','pk':'id_cliente','campos':['nombre_razon_social','nit_dpi','telefono','correo','direccion_fiscal'],'titulo':'Clientes','roles':VENTA},
 'proveedores': {'tabla':'proveedor','pk':'id_proveedor','campos':['razon_social','nit','contacto','telefono','correo','direccion'],'titulo':'Proveedores','roles':INV},
}

@app.route('/gestion/<modulo>', methods=['GET', 'POST'])
@login_required
def gestion(modulo):
    config = CRUD.get(modulo)
    if not config:
        abort(404)

    if request.method == 'POST':
        if current_user.rol not in config['roles']:
            abort(403)
        try:
            campos = config['campos']
            data = {c: request.form.get(c, '').strip() or None for c in campos}

            
            if modulo == 'categorias':
                data['nombre'] = solo_letras(
                    required('nombre'),
                    'Nombre de categoría'
                )

                if data['descripcion']:
                    data['descripcion'] = solo_letras(
                        data['descripcion'],
                        'Descripción de categoría'
                    )

            if modulo == 'proveedores' and data['contacto']:
                data['contacto'] = solo_letras(
                    data['contacto'],
                    'Contacto del proveedor'
                )

            data = validar_datos_gestion(modulo, data)
            rid = request.form.get('id', '').strip()

            if rid:
                if not rid.isdigit():
                    raise ValueError('El identificador del registro no es válido.')
                existente = one(
                    f"SELECT {config['pk']} FROM {config['tabla']} WHERE {config['pk']}=:rid",
                    rid=int(rid)
                )
                if not existente:
                    raise ValueError('El registro que intentas editar no existe.')

            # Evitar duplicados de documentos, sin impedir editar el mismo registro.
            if modulo in ('clientes', 'proveedores'):
                documento = 'nit_dpi' if modulo == 'clientes' else 'nit'
                nit = data.get(documento)
                if nit:
                    duplicado = one(
                        f"SELECT {config['pk']} FROM {config['tabla']} "
                        f"WHERE UPPER({documento})=:documento "
                        f"AND {config['pk']}<>:actual LIMIT 1",
                        documento=nit,
                        actual=int(rid) if rid else -1
                    )
                    if duplicado:
                        raise ValueError('Ya existe un registro con ese NIT o DPI.')

            if rid:
                changes = ','.join(f'{c}=:{c}' for c in campos)
                db.session.execute(
                    text(f"UPDATE {config['tabla']} SET {changes} WHERE {config['pk']}=:rid"),
                    {**data, 'rid': int(rid)}
                )
            else:
                db.session.execute(
                    text(
                        f"INSERT INTO {config['tabla']}({','.join(campos)}) "
                        f"VALUES({','.join(':'+c for c in campos)})"
                    ), data
                )

            db.session.commit()
            flash('Registro guardado correctamente.', 'success')
        except ValueError as e:
            db.session.rollback()
            flash(str(e), 'danger')
        except IntegrityError:
            db.session.rollback()
            flash('No se pudo guardar: un dato está duplicado o no es válido.', 'danger')
        except Exception:
            db.session.rollback()
            app.logger.exception('Error al guardar en gestión: %s', modulo)
            flash('No fue posible guardar el registro.', 'danger')
        return redirect(url_for('gestion', modulo=modulo))

    registros = rows(f"SELECT * FROM {config['tabla']} ORDER BY {config['pk']} DESC")
    return render_template('gestion.html', modulo=modulo, config=config, data=registros)

@app.get('/pos')
@role_required(*VENTA)
def pos():
    products=rows('''SELECT p.id_producto,p.codigo,p.codigo_barras,p.nombre,p.precio_venta,p.id_categoria,
                     vi.existencia_actual FROM producto p JOIN vista_inventario vi USING(id_producto)
                     WHERE p.estado=TRUE ORDER BY p.codigo''')
    clientes=rows('SELECT id_cliente,nombre_razon_social FROM cliente ORDER BY nombre_razon_social')
    turno=one('''SELECT * FROM turno_caja WHERE id_usuario=:uid AND estado='ABIERTO'
                 ORDER BY id_turno DESC LIMIT 1''',uid=int(current_user.id))
    return render_template('pos.html',products=products,clientes=clientes,turno=turno)

@app.post('/pos/cobrar')
@role_required(*VENTA)
def cobrar():
    """Recalcula montos en el servidor y protege stock con bloqueos de filas."""
    from flask import session
    try:
        data=request.get_json(silent=True) or {}
        cart=data.get('items',[])
        if not isinstance(cart,list) or not cart or len(cart)>80: raise ValueError('Carrito vacío o inválido.')
        metodo=data.get('metodo','')
        if metodo not in ('EFECTIVO','TARJETA','TRANSFERENCIA'): raise ValueError('Método de pago inválido.')
        cliente_id=int(data.get('cliente_id'))
        efectivo=money(data.get('recibido',0))
        turno=one('''SELECT id_turno FROM turno_caja WHERE id_usuario=:u AND estado='ABIERTO'
                     ORDER BY id_turno DESC LIMIT 1 FOR UPDATE''',u=int(current_user.id))
        if not turno: raise ValueError('Debes abrir caja para vender.')
        if not one('SELECT id_cliente FROM cliente WHERE id_cliente=:id',id=cliente_id): raise ValueError('Cliente inválido.')
        # Descartar duplicados y bloquear siempre por ID ascendente.
        cantidades={}
        descuentos={}
        for item in cart:
            pid=int(item['id']); cant=int(item['cantidad']); desc=money(item.get('descuento',0))
            if cant<=0 or pid in cantidades or desc<0: raise ValueError('Productos duplicados o cantidades inválidas.')
            cantidades[pid]=cant; descuentos[pid]=desc
        products=[]; total=Decimal('0.00')
        for pid in sorted(cantidades):
            p=one('''SELECT id_producto,nombre,precio_venta,costo FROM producto
                     WHERE id_producto=:id AND estado=TRUE FOR UPDATE''',id=pid)
            if not p: raise ValueError('Producto inactivo o inexistente.')
            saldo=value('SELECT existencia_actual FROM vista_inventario WHERE id_producto=:id',id=pid)
            cant=cantidades[pid]; desc=descuentos[pid]
            bruto=money(p['precio_venta'])*cant
            if cant>saldo: raise ValueError(f'Stock insuficiente para {p["nombre"]}: {saldo}.')
            if desc>bruto: raise ValueError('Descuento superior al importe del producto.')
            total+=bruto-desc
            products.append((p,cant,desc))
        total=money(total)
        if metodo=='EFECTIVO' and efectivo<total: raise ValueError('El efectivo recibido es insuficiente.')
        numero='PC-'+datetime.now(TZ).strftime('%Y%m%d')+'-'+secrets.token_hex(4).upper()
        vid=value('''INSERT INTO venta(id_cliente,id_usuario,id_turno,numero_ticket,metodo_pago,estado,total)
                     VALUES(:cliente,:usuario,:turno,:ticket,:metodo,'COMPLETADA',:total) RETURNING id_venta''',
                     cliente=cliente_id,usuario=int(current_user.id),turno=turno['id_turno'],ticket=numero,metodo=metodo,total=total)
        for p,cant,desc in products:
            db.session.execute(text('''INSERT INTO detalle_venta(id_venta,id_producto,cantidad,precio_unitario,
               costo_unitario_historico,descuento,tasa_iva)
               VALUES(:v,:p,:n,:precio,:costo,:desc,12)'''),{'v':vid,'p':p['id_producto'],'n':cant,
               'precio':p['precio_venta'],'costo':p['costo'],'desc':desc})
            db.session.execute(text('''INSERT INTO movimiento_inventario(id_producto,tipo,cantidad,costo_unitario,motivo)
               VALUES(:p,'SALIDA',:n,:costo,:motivo)'''),{'p':p['id_producto'],'n':cant,'costo':p['costo'],'motivo':f'Venta {numero}'})
        if metodo=='EFECTIVO':
            db.session.execute(text('''INSERT INTO movimiento_caja(id_turno,tipo,concepto,monto)
               VALUES(:turno,'INGRESO',:concepto,:monto)'''),{'turno':turno['id_turno'],'concepto':f'Venta {numero}','monto':total})
        db.session.commit()
        return jsonify(ok=True,id_venta=vid,numero=numero,total=str(total),cambio=str(efectivo-total if metodo=='EFECTIVO' else Decimal('0')))
    except (ValueError, TypeError, KeyError, InvalidOperation) as e:
        db.session.rollback(); return jsonify(ok=False,error=str(e)),400
    except Exception:
        db.session.rollback(); app.logger.exception('Cobrar venta'); return jsonify(ok=False,error='No se pudo guardar la venta.'),500

@app.get('/ventas')
@login_required
def ventas():
    filtro=request.args.get('q','').strip()
    data=rows('''SELECT v.*,c.nombre_razon_social,u.nombre AS cajero
                 FROM venta v JOIN cliente c USING(id_cliente) JOIN usuario u USING(id_usuario)
                 WHERE v.numero_ticket ILIKE :q ORDER BY v.id_venta DESC LIMIT 250''',q=f'%{filtro}%')
    return render_template('ventas.html',data=data)

@app.get('/ventas/<int:vid>/ticket')
@login_required
def ticket(vid):
    venta=one('''SELECT v.*,c.nombre_razon_social,c.nit_dpi,u.nombre AS cajero
             FROM venta v JOIN cliente c USING(id_cliente) JOIN usuario u USING(id_usuario)
             WHERE v.id_venta=:id''',id=vid)
    if not venta: abort(404)
    detalle=rows('''SELECT d.*,p.nombre FROM detalle_venta d JOIN producto p USING(id_producto)
                   WHERE d.id_venta=:id ORDER BY d.id_detalle_venta''',id=vid)
    return render_template('ticket.html',venta=venta,detalle=detalle)

@app.get('/compras')
@login_required
def compras():
    data=rows('''SELECT c.*,p.razon_social,u.nombre AS usuario
                FROM compra c JOIN proveedor p USING(id_proveedor) JOIN usuario u USING(id_usuario)
                ORDER BY c.id_compra DESC LIMIT 150''')
    return render_template('compras.html',data=data,proveedores=rows('SELECT * FROM proveedor ORDER BY razon_social'),
                           productos=rows('SELECT id_producto,codigo,nombre,costo FROM producto WHERE estado=TRUE ORDER BY codigo'))

@app.post('/compras/nueva')
@role_required(*INV)
def compra_nueva():
    try:
        proveedor=number('id_proveedor',1); producto=number('id_producto',1)
        cantidad=number('cantidad',1); costo=numeric('costo_unitario')
        metodo=required('metodo_pago')
        factura=request.form.get('numero_factura','').strip()
        p=one('SELECT id_producto FROM producto WHERE id_producto=:id FOR UPDATE',id=producto)
        if not p: raise ValueError('Producto no encontrado.')
        if not one('SELECT id_proveedor FROM proveedor WHERE id_proveedor=:id',id=proveedor): raise ValueError('Proveedor inválido.')
        total=money(costo*cantidad)
        cid=value('''INSERT INTO compra(id_proveedor,id_usuario,numero_factura,metodo_pago,total)
                     VALUES(:prov,:user,:factura,:metodo,:total) RETURNING id_compra''',
                  prov=proveedor,user=int(current_user.id),factura=factura,metodo=metodo,total=total)
        db.session.execute(text('''INSERT INTO detalle_compra(id_compra,id_producto,cantidad,costo_unitario)
                       VALUES(:compra,:producto,:cantidad,:costo)'''),{
                       'compra':cid,'producto':producto,'cantidad':cantidad,'costo':costo})
        db.session.execute(text('''INSERT INTO movimiento_inventario(id_producto,tipo,cantidad,costo_unitario,motivo)
                    VALUES(:producto,'ENTRADA',:cantidad,:costo,:motivo)'''),{
                    'producto':producto,'cantidad':cantidad,'costo':costo,'motivo':f'Compra #{cid}'})
        db.session.execute(text('UPDATE producto SET costo=:costo WHERE id_producto=:id'),{'costo':costo,'id':producto})
        db.session.commit(); flash('Compra registrada y stock actualizado.','success')
    except (ValueError,IntegrityError) as e:
        db.session.rollback(); flash(str(e) if isinstance(e,ValueError) else 'Datos de compra incorrectos.','danger')
    except Exception:
        db.session.rollback(); app.logger.exception('Compra'); flash('Error al guardar la compra.','danger')
    return redirect(url_for('compras'))

@app.route('/caja',methods=['GET','POST'])
@role_required(*VENTA)
def caja():
    if request.method=='POST':
        accion=request.form.get('accion')
        try:
            turno=one('''SELECT * FROM turno_caja WHERE id_usuario=:u AND estado='ABIERTO'
                         ORDER BY id_turno DESC LIMIT 1 FOR UPDATE''',u=int(current_user.id))
            if accion=='abrir':
                if turno: raise ValueError('Ya tienes un turno abierto.')
                db.session.execute(text('''INSERT INTO turno_caja(id_usuario,fondo_inicial,estado)
                          VALUES(:u,:fondo,'ABIERTO')'''),{'u':int(current_user.id),'fondo':numeric('fondo_inicial')})
            elif accion=='movimiento':
                if not turno: raise ValueError('Abre caja primero.')
                tipo=required('tipo')
                if tipo not in ('INGRESO','EGRESO'): raise ValueError('Tipo inválido.')
                db.session.execute(text('''INSERT INTO movimiento_caja(id_turno,tipo,concepto,monto)
                            VALUES(:turno,:tipo,:concepto,:monto)'''),{'turno':turno['id_turno'],'tipo':tipo,
                            'concepto':required('concepto'),'monto':numeric('monto',Decimal('0.01'))})
            elif accion=='cerrar':
                if not turno: raise ValueError('No tienes un turno abierto.')
                contado=numeric('efectivo_contado')
                db.session.execute(text('''UPDATE turno_caja SET cierre=NOW(),estado='CERRADO',
                                    efectivo_contado=:contado WHERE id_turno=:id'''),
                                    {'contado':contado,'id':turno['id_turno']})
            else: raise ValueError('Acción no válida.')
            db.session.commit(); flash('Operación de caja guardada.','success')
        except (ValueError,IntegrityError) as e:
            db.session.rollback(); flash(str(e) if isinstance(e,ValueError) else 'Datos incorrectos.','danger')
        except Exception:
            db.session.rollback(); app.logger.exception('Caja'); flash('No se pudo actualizar caja.','danger')
        return redirect(url_for('caja'))
    turno=one('''SELECT * FROM turno_caja WHERE id_usuario=:u AND estado='ABIERTO'
                 ORDER BY id_turno DESC LIMIT 1''',u=int(current_user.id))
    mov=[]; saldo=None
    if turno:
        mov=rows('SELECT * FROM movimiento_caja WHERE id_turno=:t ORDER BY fecha_hora DESC',t=turno['id_turno'])
        total_in=sum(m['monto'] for m in mov if m['tipo']=='INGRESO')
        total_out=sum(m['monto'] for m in mov if m['tipo']=='EGRESO')
        saldo=turno['fondo_inicial']+total_in-total_out
    ultimos=rows('''SELECT t.*,u.nombre AS usuario FROM turno_caja t JOIN usuario u USING(id_usuario)
                    ORDER BY t.id_turno DESC LIMIT 30''')
    return render_template('caja.html',turno=turno,mov=mov,saldo=saldo,ultimos=ultimos)

@app.get('/reportes')
@login_required
def reportes():
    desde=request.args.get('desde') or (date.today()-timedelta(days=29)).isoformat()
    hasta=request.args.get('hasta') or date.today().isoformat()
    try:
        d1=date.fromisoformat(desde); d2=date.fromisoformat(hasta)
        if d2<d1 or (d2-d1).days>366: raise ValueError()
    except ValueError:
        flash('Rango de fechas inválido o mayor a un año.','danger'); return redirect(url_for('reportes'))
    resumen=one('''SELECT COUNT(*) AS cantidad,COALESCE(SUM(total),0) AS total FROM venta
                   WHERE fecha_hora::date BETWEEN :inicio AND :fin AND estado='COMPLETADA' ''',inicio=d1,fin=d2)
    dias=rows('''SELECT fecha_hora::date AS fecha,SUM(total) AS total FROM venta
                WHERE fecha_hora::date BETWEEN :inicio AND :fin AND estado='COMPLETADA'
                GROUP BY fecha_hora::date ORDER BY fecha''',inicio=d1,fin=d2)
    productos=rows('''SELECT p.nombre,SUM(d.cantidad) AS cantidad,
                  SUM(d.precio_unitario*d.cantidad-d.descuento) AS total
                  FROM detalle_venta d JOIN venta v USING(id_venta) JOIN producto p USING(id_producto)
                  WHERE v.fecha_hora::date BETWEEN :inicio AND :fin AND v.estado='COMPLETADA'
                  GROUP BY p.nombre ORDER BY cantidad DESC LIMIT 15''',inicio=d1,fin=d2)
    return render_template('reportes.html',resumen=resumen,dias=dias,productos=productos,desde=desde,hasta=hasta)

@app.get('/reportes/excel')
@login_required
def reporte_excel():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='Ventas';ws.append(['Ticket','Fecha','Cliente','Cajero','Método','Total Q','Estado'])
    for v in rows('''SELECT v.numero_ticket,v.fecha_hora,c.nombre_razon_social,u.nombre AS cajero,
                       v.metodo_pago,v.total,v.estado FROM venta v JOIN cliente c USING(id_cliente)
                       JOIN usuario u USING(id_usuario) ORDER BY v.id_venta DESC LIMIT 5000'''):
        ws.append([v['numero_ticket'],str(v['fecha_hora']),v['nombre_razon_social'],v['cajero'],v['metodo_pago'],float(v['total']),v['estado']])
    bio=BytesIO();wb.save(bio);bio.seek(0)
    return send_file(bio,as_attachment=True,download_name='punto_com_ventas.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.get('/usuarios')
@role_required(*ADMIN)
def usuarios():
    return render_template('usuarios.html',data=rows('''SELECT u.id_usuario,u.nombre,u.correo,u.estado,r.nombre AS rol
      FROM usuario u JOIN rol r USING(id_rol) ORDER BY u.id_usuario'''),roles=rows('SELECT * FROM rol ORDER BY nombre'))

@app.post('/usuarios')
@role_required(*ADMIN)
def usuario_crear():
    try:
        nombre=solo_letras(required('nombre'), 'Nombre de usuario');correo=required('correo').lower();rol=number('id_rol',1);pw=required('password')
        if len(pw)<12: raise ValueError('La contraseña debe tener al menos 12 caracteres.')
        db.session.execute(text('''INSERT INTO usuario(id_rol,nombre,correo,hash_contrasena,estado)
               VALUES(:rol,:nombre,:correo,:hash,TRUE)'''),{'rol':rol,'nombre':nombre,'correo':correo,'hash':generate_password_hash(pw)})
        db.session.commit();flash('Usuario creado.','success')
    except (ValueError,IntegrityError) as e:
        db.session.rollback();flash(str(e) if isinstance(e,ValueError) else 'Correo duplicado o rol inválido.','danger')
    return redirect(url_for('usuarios'))

@app.post('/usuarios/<int:uid>/estado')
@role_required(*ADMIN)
def usuario_estado(uid):
    if uid==int(current_user.id): flash('No puedes desactivar tu propia cuenta.','danger')
    else:
        db.session.execute(text('UPDATE usuario SET estado=NOT estado WHERE id_usuario=:id'),{'id':uid});db.session.commit()
        flash('Usuario actualizado.','success')
    return redirect(url_for('usuarios'))

@app.get('/configuracion')
@login_required
def configuracion():
    return render_template('message.html',titulo='Configuración',mensaje='Nombre comercial: PUNTO COM · Moneda: Quetzal (Q) · País: Guatemala. La edición de datos fiscales e impresora estará disponible en una siguiente versión.')

@app.cli.command('crear-admin')
def crear_admin():
    """Comando CLI: flask --app app crear-admin"""
    from getpass import getpass
    nombre=input('Nombre del administrador: ').strip(); correo=input('Correo electrónico: ').strip().lower()
    password=getpass('Contraseña (mínimo 12 caracteres): ')
    if not nombre or not correo or len(password)<12: raise ValueError('Nombre, correo y contraseña >=12 son obligatorios.')
    rol=value("SELECT id_rol FROM rol WHERE nombre='Administrador'")
    if not rol: raise ValueError('Primero inserta el rol Administrador en PostgreSQL.')
    db.session.execute(text('''INSERT INTO usuario(id_rol,nombre,correo,hash_contrasena,estado)
                    VALUES(:rol,:nombre,:correo,:hash,TRUE)'''),{'rol':rol,'nombre':nombre,'correo':correo,'hash':generate_password_hash(password)})
    db.session.commit();print('Administrador registrado correctamente.')


@app.route("/reportes/ventas/excel")
@login_required
def exportar_ventas_excel():

    fecha_inicio = request.args.get("inicio", "")
    fecha_fin = request.args.get("fin", "")

    try:
        inicio = date.fromisoformat(fecha_inicio)
        fin = date.fromisoformat(fecha_fin)

        if inicio > fin:
            raise ValueError("Rango inválido")

    except ValueError:
        flash("Selecciona un rango de fechas válido.", "danger")
        return redirect(url_for("reportes"))

    fin_exclusivo = fin + timedelta(days=1)

    ventas = db.session.execute(
        text("""
            SELECT
                v.numero_ticket,
                v.fecha_hora,
                c.nombre_razon_social AS cliente,
                u.nombre AS cajero,
                v.metodo_pago,
                v.total,
                v.estado
            FROM venta v
            INNER JOIN cliente c
                ON v.id_cliente = c.id_cliente
            INNER JOIN usuario u
                ON v.id_usuario = u.id_usuario
            WHERE v.fecha_hora >= :inicio
              AND v.fecha_hora < :fin
            ORDER BY v.fecha_hora DESC
        """),
        {
            "inicio": inicio,
            "fin": fin_exclusivo
        }
    ).mappings().all()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas PUNTO COM"

    encabezados = [
        "Ticket",
        "Fecha",
        "Cliente",
        "Cajero",
        "Método de pago",
        "Total Q",
        "Estado"
    ]

    hoja.append(encabezados)

    for celda in hoja[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill(
            "solid", fgColor="1D4ED8"
        )

    for venta in ventas:
        hoja.append([
            venta["numero_ticket"],
            venta["fecha_hora"],
            venta["cliente"],
            venta["cajero"],
            venta["metodo_pago"],
            float(venta["total"]),
            venta["estado"]
        ])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for fila in hoja.iter_rows(min_row=2):
        fila[1].number_format = "DD/MM/YYYY HH:MM"
        fila[5].number_format = '"Q" #,##0.00'

    anchos = {
        "A": 20, "B": 22, "C": 30,
        "D": 25, "E": 20, "F": 18,
        "G": 18
    }

    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name="reporte_ventas_punto_com.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

if __name__=='__main__':
    app.run(debug=True)
